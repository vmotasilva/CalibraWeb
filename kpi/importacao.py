"""Modelos de planilha (diário, semanal, mensal) e importação de valores e motivos dos indicadores."""
import datetime
import io
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kpi import calculo
from kpi.models import Indicador, Medicao, Motivo
from kpi.views import _normalizar, _rotulo

CABECALHO_PERIODO = {'diaria': 'Dia', 'semanal': 'Semana', 'mensal': 'Mês'}
NOMES = {'diaria': 'Diario', 'semanal': 'Semanal', 'mensal': 'Mensal'}
EXEMPLOS = {
    'diaria': ('15/10/2026', 'Dia do lançamento. Ex.: 15/10/2026'),
    'semanal': ('12/10/2026', 'Qualquer dia da semana (vale segunda a domingo). Ex.: 12/10/2026'),
    'mensal': ('10/2026', 'Mês no formato mm/aaaa (ou qualquer data do mês). Ex.: 10/2026'),
}
ABA_VALORES, ABA_MOTIVOS = 'Valores', 'Motivos'


TIPOS = {'valores': 'Gráficos', 'motivos': 'Tabelas'}
ABA_POR_TIPO = {'valores': ABA_VALORES, 'motivos': ABA_MOTIVOS}


def gerar_modelo(gran, ind_especifico=None, tipo='valores'):
    """Cria o .xlsx modelo. tipo='valores' alimenta os gráficos; tipo='motivos' alimenta as tabelas Top motivos.
    Com `ind_especifico`, inclui os campos e setores definidos nesse indicador."""
    periodo = CABECALHO_PERIODO[gran]
    exemplo, dica = EXEMPLOS[gran]
    nomes = [ind_especifico.nome] if ind_especifico else list(Indicador.objects.values_list('nome', flat=True))
    nome_ex = nomes[0] if nomes else 'Nome do indicador'
    campos = calculo.lista_campos(ind_especifico) if ind_especifico else []
    formula = ind_especifico.formula if ind_especifico else ''
    setores = list(ind_especifico.setores.all()) if ind_especifico else []
    col_setor = ['Setor'] if setores else []
    vazio_setor = [None] if setores else []
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_POR_TIPO[tipo]
    if tipo == 'valores':
        cab_valor = 'Valor (opcional: calculado pela fórmula)' if formula else 'Valor'
        ws.append(['Indicador', periodo] + col_setor + campos + [cab_valor, 'Observação'])
        ws.append([nome_ex, exemplo] + vazio_setor + [1] * len(campos)
                  + [None if formula else 1.5, 'exemplo (apague esta linha)'])
        instrucoes = [
            ['Uma linha por período (colunas: Indicador, ' + periodo + (', Setor' if setores else '')
             + (', ' + ', '.join(campos) if campos else '') + ', Valor, Observação).'],
            *([['Campos deste indicador: ' + ', '.join(campos) + '.']] if campos else []),
            *([['Fórmula do valor: ' + formula + ' — se a coluna Valor ficar vazia, o valor é calculado a partir dos campos.']] if formula else []),
            ['Valor já existente no mesmo período é atualizado.'],
        ]
    else:
        ws.append(['Indicador', periodo, 'Motivo', 'Quantidade'] + col_setor)
        ws.append([nome_ex, exemplo, 'Exemplo de motivo', 3] + vazio_setor)
        ws.append([nome_ex, exemplo, 'Outro motivo (mesmo período)', 1] + vazio_setor)
        instrucoes = [
            ['Uma linha por motivo (colunas: Indicador, ' + periodo + ', Motivo, Quantidade' + (', Setor' if setores else '') + ').'],
            ['Repita Indicador e ' + periodo + ' para vários motivos no mesmo período.'],
            ['Os motivos de um período presente na planilha substituem os motivos atuais desse período.'],
        ]
    info = wb.create_sheet('Instruções')
    for linha in [
        [f'Modelo de importação de {TIPOS[tipo].upper()} — {NOMES[gran]}'],
        [],
        *instrucoes,
        *([['Setores deste indicador: ' + ', '.join(x.nome + (f' (meta {x.meta})' if x.meta is not None else '') for x in setores)
            + '. Deixe a coluna Setor vazia para o acompanhamento geral.']] if setores else []),
        [dica],
        ['O nome do indicador deve ser igual ao cadastrado (veja a aba Indicadores).'],
        ['Apague as linhas de exemplo antes de importar. Use ponto ou vírgula como separador decimal.'],
    ]:
        info.append(linha)
    ind = wb.create_sheet('Indicadores')
    ind.append(['Indicadores cadastrados', 'Meta'])
    for i in ([ind_especifico] if ind_especifico else Indicador.objects.all()):
        ind.append([i.nome, float(i.meta)])
    cor = PatternFill('solid', fgColor='14284B')
    for folha in (ws, ind):
        for c in folha[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = cor
            c.alignment = Alignment(horizontal='center')
        for n in range(1, folha.max_column + 1):
            folha.column_dimensions[get_column_letter(n)].width = 34 if n == 1 else 22
    info.column_dimensions['A'].width = 120
    info['A1'].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _data(v):
    """Aceita date/datetime, dd/mm/aaaa, mm/aaaa, aaaa-mm e aaaa-mm-dd."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    t = str(v or '').strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d/%m/%y', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    m = re.fullmatch(r'(\d{1,2})/(\d{4})', t)
    if m:
        return datetime.date(int(m.group(2)), int(m.group(1)), 1)
    m = re.fullmatch(r'(\d{4})-(\d{1,2})', t)
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), 1)
    raise ValueError(f'data inválida: {t or "vazia"}')


def _numero(v):
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    try:
        t = str(v).strip().replace('%', '').replace(' ', '')
        if ',' in t:
            t = t.replace('.', '').replace(',', '.')
        return Decimal(t)
    except (InvalidOperation, AttributeError):
        raise ValueError(f'número inválido: {v}')


def _linhas(ws):
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and any(c not in (None, '') for c in row):
            yield n, row


def consolidar_motivos_diarios(chaves, usuario):
    """Compõe as tabelas de motivos semanal e mensal a partir das diárias (soma por descrição).
    chaves: conjunto de (indicador, setor_id, data_diária) alterados na importação.
    Só recompõe períodos sem motivos lançados manualmente."""
    periodos = {(g, ind.pk, setor_id, _normalizar(g, d)): ind
                for ind, setor_id, d in chaves for g in ('semanal', 'mensal') if ind.permite_motivos(g)}
    for (g, _pk, setor_id, periodo), ind in periodos.items():
        existentes = Motivo.objects.filter(indicador=ind, setor_id=setor_id, granularidade=g, data=periodo)
        if existentes.exclude(consolidado=True).exists():
            continue
        ini, fim = (periodo, periodo + datetime.timedelta(days=6)) if g == 'semanal' else (periodo.replace(day=1), periodo)
        soma = defaultdict(Decimal)
        for desc, q in Motivo.objects.filter(indicador=ind, setor_id=setor_id, granularidade='diaria',
                                             data__range=(ini, fim)).values_list('descricao', 'quantidade'):
            soma[desc] += q
        existentes.delete()
        Motivo.objects.bulk_create([
            Motivo(indicador=ind, setor_id=setor_id, granularidade=g, data=periodo, descricao=d, quantidade=q,
                   consolidado=True, registrado_por=usuario) for d, q in soma.items()])


def importar(arquivo, gran, usuario, tipo='valores'):
    """Importa a planilha do tipo informado (valores = gráficos; motivos = tabelas).
    Retorna dict com contadores e a lista de erros (linha, aba, mensagem)."""
    try:
        wb = load_workbook(arquivo, data_only=True)
    except Exception:
        return {'valores': 0, 'motivos': 0, 'erros': [('—', '—', 'Arquivo inválido: envie um .xlsx gerado pelo modelo.')]}
    aba = ABA_POR_TIPO[tipo]
    if aba not in wb.sheetnames:
        outro = TIPOS['motivos' if tipo == 'valores' else 'valores']
        return {'valores': 0, 'motivos': 0, 'tipo': tipo, 'rotulo': CABECALHO_PERIODO[gran], 'erros': [(
            '—', '—', f'Esta planilha não é um modelo de {TIPOS[tipo].upper()} (falta a aba "{aba}"). '
                      f'Parece ser de {outro.upper()}: use a importação correspondente.')]}
    indicadores = {i.nome.strip().lower(): i for i in Indicador.objects.prefetch_related('setores')}
    erros, valores, motivos_n = [], 0, 0
    alt_motivos = set()  # (indicador, setor_id, data) diários alterados

    def achar_setor(aba, n, ind, nome_setor):
        """Retorna (setor, valido); vazio = acompanhamento geral."""
        nome_setor = str(nome_setor or '').strip()
        if not nome_setor:
            return None, True
        for s in ind.setores.all():
            if s.nome.strip().lower() == nome_setor.lower():
                return s, True
        erros.append((n, aba, f'Setor "{nome_setor}" não cadastrado em "{ind.nome}"'))
        return None, False

    def achar(aba, n, nome, periodo_bruto):
        ind = indicadores.get(str(nome or '').strip().lower())
        if not ind:
            erros.append((n, aba, f'Indicador não encontrado: "{nome}"'))
            return None, None
        if aba == ABA_MOTIVOS and not ind.permite_motivos(gran):
            erros.append((n, aba, f'"{ind.nome}" não tem a tabela de motivos {CABECALHO_PERIODO[gran].lower()} habilitada'))
            return None, None
        if aba == ABA_VALORES and not ind.permite_valor(gran):
            erros.append((n, aba, f'"{ind.nome}" não tem o gráfico {CABECALHO_PERIODO[gran].lower()} habilitado'))
            return None, None
        try:
            return ind, _normalizar(gran, _data(periodo_bruto))
        except ValueError as e:
            erros.append((n, aba, str(e)))
            return None, None

    with transaction.atomic():
        if tipo == 'valores':
            planilha = wb[ABA_VALORES]
            cab = [str(c or '').strip() for c in next(planilha.iter_rows(min_row=1, max_row=1, values_only=True), [])]
            col_val = next((i for i, c in enumerate(cab) if c.lower().startswith('valor')), 2)
            col_obs = next((i for i, c in enumerate(cab) if c.lower().startswith('observa')), None)
            col_setor = next((i for i, c in enumerate(cab) if c.lower() == 'setor'), None)
            for n, row in _linhas(planilha):
                row = list(row) + [None] * (len(cab) + 1)
                ind, periodo = achar(ABA_VALORES, n, row[0], row[1])
                if not ind:
                    continue
                setor, setor_ok = achar_setor(ABA_VALORES, n, ind, row[col_setor] if col_setor is not None else None)
                if not setor_ok:
                    continue
                dados, falha = {}, False
                for nome_campo in calculo.lista_campos(ind):
                    col = next((i for i, c in enumerate(cab) if c.lower() == nome_campo.lower()), None)
                    if col is None or row[col] in (None, ''):
                        continue
                    try:
                        dados[nome_campo] = _numero(row[col])
                    except ValueError as e:
                        erros.append((n, ABA_VALORES, f'{nome_campo}: {e}'))
                        falha = True
                if falha:
                    continue
                try:
                    if row[col_val] not in (None, ''):
                        valor = _numero(row[col_val])
                    elif ind.formula:
                        valor = calculo.calcular(ind.formula, dados)
                    else:
                        raise ValueError('valor vazio')
                except (ValueError, ZeroDivisionError) as e:
                    erros.append((n, ABA_VALORES, str(e)))
                    continue
                obs = row[col_obs] if col_obs is not None else ''
                Medicao.objects.update_or_create(
                    indicador=ind, setor=setor, periodo=periodo, granularidade=gran,
                    defaults={'valor': valor, 'observacao': str(obs or '')[:255],
                              'dados': {k: str(v) for k, v in dados.items()}, 'registrado_por': usuario})
                valores += 1
        if tipo == 'motivos':
            grupos = {}
            cab_m = [str(c or '').strip().lower() for c in next(wb[ABA_MOTIVOS].iter_rows(min_row=1, max_row=1, values_only=True), [])]
            col_setor_m = cab_m.index('setor') if 'setor' in cab_m else None
            for n, row in _linhas(wb[ABA_MOTIVOS]):
                row = list(row) + [None] * (len(cab_m) + 4)
                nome, per, desc, qtd = row[:4]
                ind, periodo = achar(ABA_MOTIVOS, n, nome, per)
                if not ind:
                    continue
                setor, setor_ok = achar_setor(ABA_MOTIVOS, n, ind, row[col_setor_m] if col_setor_m is not None else None)
                if not setor_ok:
                    continue
                desc = str(desc or '').strip()[:200]
                if not desc:
                    erros.append((n, ABA_MOTIVOS, 'Motivo vazio'))
                    continue
                try:
                    q = _numero(qtd) if qtd not in (None, '') else Decimal(1)
                except ValueError as e:
                    erros.append((n, ABA_MOTIVOS, str(e)))
                    continue
                grupos.setdefault((ind.pk, setor.pk if setor else None, periodo), []).append((ind, setor, desc, q))
            for (ind_pk, setor_pk, periodo), itens in grupos.items():
                Motivo.objects.filter(indicador_id=ind_pk, setor_id=setor_pk, granularidade=gran, data=periodo).delete()
                Motivo.objects.bulk_create([
                    Motivo(indicador=ind, setor=setor, granularidade=gran, data=periodo, descricao=d, quantidade=q, registrado_por=usuario)
                    for ind, setor, d, q in itens])
                motivos_n += len(itens)
                if gran == 'diaria':
                    alt_motivos.add((itens[0][0], setor_pk, periodo))
        if gran == 'diaria':
            consolidar_motivos_diarios(alt_motivos, usuario)
    return {'valores': valores, 'motivos': motivos_n, 'erros': erros, 'tipo': tipo, 'rotulo': CABECALHO_PERIODO[gran]}

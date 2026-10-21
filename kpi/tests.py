from datetime import date

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from kpi.models import Indicador, Medicao, Motivo


class KpiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('kpi', 'k@x.com', 'x')
        self.client.force_login(self.user)
        self.ind = Indicador.objects.create(nome='SR', meta=95, exibir_semanal=True, exibir_diario=True, motivos_diario=True, motivos_semanal=True, motivos_mensal=True)

    def salvar(self, gran, periodo, valor):
        return self.client.post(reverse('kpi:medicao_salvar', args=[self.ind.pk]),
                                {'granularidade': gran, 'periodo': periodo, 'valor': valor, 'observacao': ''})

    def test_cria_indicador(self):
        r = self.client.post(reverse('kpi:indicador_novo'), {
            'nome': 'Retrabalho', 'unidade': '%', 'meta': '4', 'sentido': 'menor',
            'periodicidade': 'mensal', 'ativo': 'on', 'exibir_mensal': 'on'})
        self.assertEqual(r.status_code, 302)
        novo = Indicador.objects.get(nome='Retrabalho')
        self.assertTrue(novo.exibir_mensal)
        self.assertFalse(novo.exibir_semanal or novo.motivos_diario)

    def test_normaliza_e_atualiza(self):
        self.salvar('mensal', '2026-10-05', '90')
        self.salvar('mensal', '2026-10-20', '96')
        m = Medicao.objects.get(granularidade='mensal')
        self.assertEqual(m.periodo, date(2026, 10, 31))
        self.assertEqual(str(m.valor), '96.00')
        self.salvar('semanal', '2026-10-15', '91')  # quinta -> segunda 12/10
        self.salvar('semanal', '2026-10-13', '92')
        s = Medicao.objects.get(granularidade='semanal')
        self.assertEqual(s.periodo, date(2026, 10, 12))
        self.salvar('diaria', '2026-10-15', '93')
        self.assertEqual(Medicao.objects.count(), 3)

    def test_sem_acompanhamento_bloqueia_semanal(self):
        self.ind.exibir_semanal = False
        self.ind.save()
        self.salvar('semanal', '2026-10-15', '91')
        self.assertEqual(Medicao.objects.count(), 0)

    def test_motivos_por_periodo(self):
        self.salvar('mensal', '2026-10-05', '90')
        url = reverse('kpi:motivo_salvar', args=[self.ind.pk])
        hoje = date.today().isoformat()
        self.client.post(url, {'granularidade': 'diaria', 'data': hoje,
                               'descricao': ['Quebra', 'quebra ', 'Trinca', ''], 'quantidade': ['3', '2', '1', '5']})
        self.assertEqual(Motivo.objects.count(), 3)
        r = self.client.get(reverse('kpi:painel') + '?ind=%d&mot=0&p_diaria=acum' % self.ind.pk)
        self.assertEqual(r.status_code, 200)
        diario = r.context['abas'][0]['motivos_tabs'][0]
        self.assertEqual(diario['gran'], 'diaria')
        self.assertEqual([t['descricao'] for t in diario['top']], ['Quebra', 'Trinca'])
        self.assertEqual(float(diario['top'][0]['total']), 5.0)
        # regravar o mesmo período substitui os motivos
        self.client.post(url, {'granularidade': 'diaria', 'data': hoje, 'descricao': ['Editado'], 'quantidade': ['9']})
        self.assertEqual(list(Motivo.objects.values_list('descricao', flat=True)), ['Editado'])
        # semana normaliza para segunda-feira; mês para o último dia
        self.client.post(url, {'granularidade': 'semanal', 'data': '2026-10-15', 'descricao': ['A'], 'quantidade': ['1']})
        self.client.post(url, {'granularidade': 'mensal', 'data': '2026-10-15', 'descricao': ['B'], 'quantidade': ['1']})
        self.assertEqual(Motivo.objects.get(granularidade='semanal').data, date(2026, 10, 12))
        self.assertEqual(Motivo.objects.get(granularidade='mensal').data, date(2026, 10, 31))
        # enviar vazio apaga o período
        self.client.post(url, {'granularidade': 'diaria', 'data': hoje, 'descricao': [''], 'quantidade': ['1']})
        self.assertEqual(Motivo.objects.filter(granularidade='diaria').count(), 0)
        self.client.post(reverse('kpi:medicao_excluir', args=[Medicao.objects.get().pk]))
        self.assertEqual(Medicao.objects.count(), 0)
        self.assertEqual(self.client.get(reverse('kpi:indicadores')).status_code, 200)

    def test_ordem_das_abas(self):
        b = Indicador.objects.create(nome='B', meta=1)
        c = Indicador.objects.create(nome='C', meta=1)
        nomes = lambda: [x.nome for x in Indicador.objects.all()]
        self.client.post(reverse('kpi:indicador_mover', args=[c.pk]), {'dir': 'up'})
        self.client.post(reverse('kpi:indicador_mover', args=[c.pk]), {'dir': 'up'})
        self.assertEqual(nomes()[0], 'C')
        r = self.client.get(reverse('kpi:painel'))
        self.assertEqual(r.context['abas'][0]['ind'].nome, 'C')
        self.assertEqual(self.client.get(reverse('kpi:indicadores')).status_code, 200)

    def test_importacao_diaria_de_motivos_compoe_semanal_e_mensal(self):
        import io
        from openpyxl import Workbook

        def enviar(linhas):
            wb = Workbook()
            ws = wb.active
            ws.title = 'Motivos'
            ws.append(['Indicador', 'Dia', 'Motivo', 'Quantidade'])
            for l in linhas:
                ws.append(l)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            buf.name = 'x.xlsx'
            return self.client.post(reverse('kpi:importar'), {'tipo': 'motivos', 'granularidade': 'diaria', 'arquivo': buf})

        enviar([['SR', '13/10/2026', 'Quebra', 3], ['SR', '14/10/2026', 'Quebra', 2], ['SR', '14/10/2026', 'Trinca', 1]])
        sem = {m.descricao: m.quantidade for m in Motivo.objects.filter(granularidade='semanal', data=date(2026, 10, 12))}
        mes = {m.descricao: m.quantidade for m in Motivo.objects.filter(granularidade='mensal', data=date(2026, 10, 31))}
        self.assertEqual(sem, {'Quebra': 5, 'Trinca': 1})
        self.assertEqual(mes, sem)
        # nova importação diária atualiza a composição; gráficos (valores) não são afetados
        enviar([['SR', '15/10/2026', 'Quebra', 4]])
        sem = Motivo.objects.get(granularidade='semanal', data=date(2026, 10, 12), descricao='Quebra')
        self.assertEqual(sem.quantidade, 9)
        self.assertFalse(Medicao.objects.filter(granularidade__in=('semanal', 'mensal')).exists())
        # lançamento manual no período não é sobrescrito
        Motivo.objects.filter(granularidade='mensal').delete()
        Motivo.objects.create(indicador=self.ind, granularidade='mensal', data=date(2026, 10, 31), descricao='Manual', quantidade=7)
        enviar([['SR', '16/10/2026', 'Quebra', 1]])
        self.assertEqual(list(Motivo.objects.filter(granularidade='mensal').values_list('descricao', flat=True)), ['Manual'])

    def test_memoria_de_calculo_com_var_e_consolidacao_automatica(self):
        from kpi import calculo
        ind = Indicador.objects.create(nome='Quebras', meta=5, sentido='menor', exibir_semanal=True, exibir_diario=True,
                                       campos='Lentes perdidas\nLentes produzidas', formula='')
        # "Var 01 / Var 02" vira [Lentes perdidas] / [Lentes produzidas]
        self.assertEqual(calculo.resolver_aliases('(Var 01 / var2) * 100', calculo.lista_campos(ind)),
                         '([Lentes perdidas] / [Lentes produzidas]) * 100')
        with self.assertRaises(ValueError):
            calculo.resolver_aliases('Var 03 / 2', calculo.lista_campos(ind))
        r = self.client.post(reverse('kpi:indicador_editar', args=[ind.pk]), {
            'nome': 'Quebras', 'meta': '5', 'sentido': 'menor', 'periodicidade': 'mensal', 'ativo': 'on',
            'exibir_mensal': 'on', 'exibir_semanal': 'on', 'exibir_diario': 'on',
            'campos': 'Lentes perdidas\nLentes produzidas', 'formula': '(Var 01 / Var 02) * 100'})
        self.assertEqual(r.status_code, 302)
        ind.refresh_from_db()
        self.assertEqual(ind.formula, '([Lentes perdidas] / [Lentes produzidas]) * 100')

        def lancar(dia, perd, prod):
            self.client.post(reverse('kpi:medicao_salvar', args=[ind.pk]), {
                'granularidade': 'diaria', 'periodo': dia, 'valor': '', 'campo_0': perd, 'campo_1': prod, 'observacao': ''})
        lancar('2026-10-13', '2', '100')
        lancar('2026-10-14', '4', '300')
        dia = Medicao.objects.get(indicador=ind, granularidade='diaria', periodo=date(2026, 10, 14))
        self.assertEqual(str(dia.valor), '1.33')
        sem = Medicao.objects.get(indicador=ind, granularidade='semanal', periodo=date(2026, 10, 12))
        mes = Medicao.objects.get(indicador=ind, granularidade='mensal', periodo=date(2026, 10, 31))
        self.assertEqual(str(sem.valor), '1.50')  # (2+4)/(100+300)*100, não a média dos dias
        self.assertEqual(str(mes.valor), '1.50')
        # editar um dia recalcula; lançamento manual da semana não é sobrescrito
        lancar('2026-10-14', '8', '300')
        self.assertEqual(str(Medicao.objects.get(pk=sem.pk).valor), '2.50')
        Medicao.objects.filter(pk=sem.pk).update(observacao='manual')
        lancar('2026-10-15', '1', '100')
        self.assertEqual(Medicao.objects.get(pk=sem.pk).observacao, 'manual')
        self.assertEqual(str(Medicao.objects.get(pk=mes.pk).valor), '2.20')
        # excluir todos os dias remove o mês consolidado
        for m in Medicao.objects.filter(indicador=ind, granularidade='diaria'):
            self.client.post(reverse('kpi:medicao_excluir', args=[m.pk]))
        self.assertFalse(Medicao.objects.filter(pk=mes.pk).exists())

    def test_consolidacao_sobrescreve_legado_sem_observacao_e_ignora_periodo_incompleto(self):
        ind = Indicador.objects.create(nome='Q3', meta=5, sentido='menor', exibir_semanal=True, exibir_diario=True,
                                       campos='P\nQ', formula='[P] / [Q] * 100')
        url = reverse('kpi:medicao_salvar', args=[ind.pk])
        # semana 12/10 já tem valor importado antigo (sem observação) e um dia antigo sem variáveis
        Medicao.objects.create(indicador=ind, granularidade='semanal', periodo=date(2026, 10, 12), valor=9)
        Medicao.objects.create(indicador=ind, granularidade='diaria', periodo=date(2026, 10, 13), valor=3)
        self.client.post(url, {'granularidade': 'diaria', 'periodo': '2026-10-14', 'valor': '', 'campo_0': '1', 'campo_1': '100', 'observacao': ''})
        self.assertEqual(str(Medicao.objects.get(indicador=ind, granularidade='semanal').valor), '9.00')  # incompleto: intacto
        # semana seguinte só com dias com variáveis: o valor antigo sem observação é substituído
        Medicao.objects.create(indicador=ind, granularidade='semanal', periodo=date(2026, 10, 19), valor=9)
        self.client.post(url, {'granularidade': 'diaria', 'periodo': '2026-10-20', 'valor': '', 'campo_0': '2', 'campo_1': '200', 'observacao': ''})
        self.assertEqual(str(Medicao.objects.get(indicador=ind, granularidade='semanal', periodo=date(2026, 10, 19)).valor), '1.00')

    def test_colagem_varios_destinos_com_operacao_e_fator(self):
        from kpi.models import ModeloColagem
        hmc = Indicador.objects.create(nome='HMC', meta=5, sentido='menor', exibir_diario=True,
                                       campos='Perdidas\nProduzidas', formula='[Perdidas] / [Produzidas] * 100')
        qb = Indicador.objects.create(nome='Qb', meta=5, sentido='menor', exibir_diario=True,
                                      campos='Perdidas\nProduzidas', formula='[Perdidas] / [Produzidas] * 100')
        modelo = ModeloColagem.objects.create(nome='Multi', linhas='\n'.join([
            'HC boas || HMC | Produzidas | | + | 1 || Qb | Produzidas | | + | 1',
            'AR boas || HMC | Produzidas | | + | 1',
            'HC quebradas || HMC | Perdidas | | + | 1 || Qb | Perdidas | | - | 1',
            'AR quebradas || HMC | Perdidas | | + | 1',
            'Ajuste || Qb | Perdidas | | + | 2',
            'Solta',
        ]))
        r = self.client.post(reverse('kpi:colagem'), {
            'modelo': modelo.pk, 'granularidade': 'diaria', 'data': '2026-10-14', 'confirmar': '1',
            'texto': '1000\n500\n20\n10\n5\n99'})
        self.assertFalse(r.context.get('erros'))
        h = Medicao.objects.get(indicador=hmc, granularidade='diaria')
        self.assertEqual(h.dados, {'Produzidas': '1500', 'Perdidas': '30'})  # soma das posições
        self.assertEqual(str(h.valor), '2.00')
        q = Medicao.objects.get(indicador=qb, granularidade='diaria')
        self.assertEqual(q.dados, {'Produzidas': '1000', 'Perdidas': '-10'})  # -20 + (5 × 2)

    def test_migracao_vincula_kpi_diario_sem_tocar_no_que_ja_foi_vinculado(self):
        import importlib
        from django.apps import apps
        from kpi import colagem
        from kpi.models import ModeloColagem, SetorIndicador
        mig = importlib.import_module('kpi.migrations.0013_vincular_kpi_diario')
        qb = Indicador.objects.create(nome='Quebras', meta=5, sentido='menor', campos='Motivo | tabela\nLentes Perdidas\nLentes Produzidas',
                                      formula='[Lentes Perdidas] / [Lentes Produzidas] * 100', exibir_diario=True)
        SetorIndicador.objects.create(indicador=qb, nome='HMC')
        ModeloColagem.objects.filter(nome='KPI Diário').update(linhas='\n'.join([
            'HC - LENTES BOAS',  # sem vínculo: recebe
            'HC - LENTES QUEBRADAS | Quebras | Lentes Perdidas | HMC',  # já vinculado à mão: intacto
            'SURFACAGEM - LENTES BOAS',  # setor Surfaçagem não existe: segue sem vínculo
        ]))
        mig.forwards(apps, None)
        linhas, erros = colagem.interpretar_modelo(ModeloColagem.objects.get(nome='KPI Diário').linhas)
        self.assertFalse(erros)
        self.assertEqual(linhas[0]['destino'], '+ Quebras · Lentes Produzidas · HMC')
        self.assertEqual(linhas[1]['destino'], '+ Quebras · Lentes Perdidas · HMC')
        self.assertEqual(linhas[2]['destinos'], [])

    def test_migracao_service_rate_define_variaveis_e_acrescenta_vinculos(self):
        import importlib
        from django.apps import apps
        from kpi import colagem
        from kpi.models import ModeloColagem
        mig = importlib.import_module('kpi.migrations.0014_service_rate_variaveis_e_vinculos')
        Indicador.objects.create(nome='Late Orders', meta=1, sentido='menor', exibir_diario=True,
                                 campos='Lentes enviadas', formula='')
        sr = Indicador.objects.create(nome='Service Rate', meta=95, exibir_diario=True)
        ModeloColagem.objects.filter(nome='KPI Diário').update(linhas='\n'.join([
            'TOTAL LENTES - ENVIADAS NO PRAZO',
            'TOTAL LENTES - ENVIADAS | Late Orders | Lentes enviadas',  # já com vínculo no formato simples
            'OUTRA',
        ]))
        mig.forwards(apps, None)
        mig.forwards(apps, None)  # idempotente
        sr.refresh_from_db()
        self.assertEqual(sr.formula, '[Lentes enviadas no prazo] / [Lentes enviadas] * 100')
        linhas, erros = colagem.interpretar_modelo(ModeloColagem.objects.get(nome='KPI Diário').linhas)
        self.assertFalse(erros)
        self.assertEqual(linhas[0]['destino'], '+ Service Rate · Lentes enviadas no prazo')
        self.assertEqual(linhas[1]['destino'], '+ Late Orders · Lentes enviadas; + Service Rate · Lentes enviadas')
        self.assertEqual(linhas[2]['destinos'], [])

    def test_colagem_linhas_em_branco_ocupam_a_posicao(self):
        from kpi.models import ModeloColagem
        ind = Indicador.objects.create(nome='Qz', meta=5, sentido='menor', exibir_diario=True,
                                       campos='P\nQ', formula='[P] / [Q] * 100')
        modelo = ModeloColagem.objects.create(nome='Brancos', linhas='\n'.join([
            'A', 'B || Qz | P | | + | 1', 'C', 'D || Qz | Q | | + | 1', 'E']))
        url = reverse('kpi:colagem')
        base = {'modelo': modelo.pk, 'granularidade': 'diaria', 'data': '2026-10-14', 'confirmar': '1'}
        # 5 posições: a 3ª e a 5ª vieram em branco e continuam ocupando o lugar delas
        r = self.client.post(url, dict(base, texto='10\n5\n\n200\n '))
        self.assertFalse(r.context.get('erros'))
        m = Medicao.objects.get(indicador=ind, granularidade='diaria')
        self.assertEqual(m.dados, {'P': '5', 'Q': '200'})  # B e D, sem deslocar as posições
        self.assertEqual(str(m.valor), '2.50')
        Medicao.objects.all().delete()
        # linhas a mais no fim: avisa e só segue se o usuário confirmar
        r = self.client.post(url, dict(base, texto='10\n5\n\n200\n \n '))
        self.assertEqual(r.context['excesso']['n'], 1)
        self.assertTrue(r.context['excesso']['vazias'])
        self.assertFalse(Medicao.objects.exists())
        r = self.client.post(url, dict(base, texto='10\n5\n\n200\n \n ', excesso_ok='1'))
        self.assertFalse(r.context.get('excesso'))
        self.assertEqual(str(Medicao.objects.get(indicador=ind, granularidade='diaria').valor), '2.50')
        # sobra com números também avisa (variáveis que o modelo não conhece)
        r = self.client.post(url, dict(base, texto='10\n5\n\n200\n\n77'))
        self.assertFalse(r.context['excesso']['vazias'])
        # faltando linhas continua sendo erro
        r = self.client.post(url, dict(base, texto='10\n5'))
        self.assertTrue(r.context['erros'])

    def test_causas_e_analises_com_rascunho_dos_registros_dos_dias(self):
        ind = Indicador.objects.create(nome='Qa', meta=2, sentido='menor', exibir_semanal=True, exibir_diario=True,
                                       campos='P\nQ', formula='[P] / [Q] * 100')
        for dia, v in (('2026-10-13', 3), ('2026-10-14', 1), ('2026-10-15', 4)):
            Medicao.objects.create(indicador=ind, granularidade='diaria', periodo=date.fromisoformat(dia), valor=v)
        Motivo.objects.create(indicador=ind, granularidade='diaria', data=date(2026, 10, 13), descricao='Risco de garra', quantidade=3)
        Motivo.objects.create(indicador=ind, granularidade='diaria', data=date(2026, 10, 15), descricao='Risco de garra', quantidade=2)
        Motivo.objects.create(indicador=ind, granularidade='semanal', data=date(2026, 10, 12), descricao='Duplicado consolidado',
                              quantidade=9, consolidado=True)
        salvar = reverse('kpi:analise_salvar', args=[ind.pk])
        dados = reverse('kpi:analise_dados', args=[ind.pk])
        r = self.client.post(salvar, {'gran': 'diaria', 'data': '2026-10-13', 'causa': 'Troca de turno', 'analise': 'Treinar'})
        self.assertTrue(r.json()['ok'])
        # a semana junta os dias fora da meta, as causas já registradas nos dias e os motivos (sem duplicar os consolidados)
        j = self.client.get(dados, {'gran': 'semanal', 'data': '2026-10-15'}).json()
        self.assertEqual(j['data'], '2026-10-12')  # normalizada para a segunda-feira
        self.assertIsNone(j['existente'])
        self.assertIn('13/10 (3)', j['rascunho'])
        self.assertIn('15/10 (4)', j['rascunho'])
        self.assertNotIn('14/10 (1)', j['rascunho'])
        self.assertIn('- 13/10: Troca de turno Treinar', j['rascunho'])
        self.assertIn('Risco de garra (5)', j['rascunho'])
        self.assertNotIn('Duplicado', j['rascunho'])
        # o mês também recolhe as análises das semanas
        self.client.post(salvar, {'gran': 'semanal', 'data': '2026-10-12', 'causa': 'Semana ruim'})
        m = self.client.get(dados, {'gran': 'mensal', 'data': '2026-10-20'}).json()
        self.assertIn('semana de 12/10: Semana ruim', m['rascunho'])
        # existente volta preenchido; vazio remove; painel entrega o mapa das análises
        self.assertEqual(self.client.get(dados, {'gran': 'diaria', 'data': '2026-10-13'}).json()['existente']['causa'], 'Troca de turno')
        r = self.client.get(reverse('kpi:painel') + '?ind=%d' % ind.pk)
        self.assertIn('2026-10-13', r.context['analises_js'][str(ind.pk)]['diaria'])
        self.client.post(salvar, {'gran': 'diaria', 'data': '2026-10-13', 'causa': '', 'analise': ''})
        self.assertFalse(Medicao.objects.none().exists() or ind.analises.filter(granularidade='diaria').exists())
        self.assertEqual(self.client.post(salvar, {'gran': 'x', 'data': '2026-10-13'}).status_code, 400)

    def test_variavel_de_tabela_fica_fora_do_calculo(self):
        from kpi import calculo
        ind = Indicador.objects.create(nome='Q2', meta=5, campos='Motivo | tabela\nPerdidas\nProduzidas')
        self.assertEqual(calculo.lista_campos(ind), ['Perdidas', 'Produzidas'])
        self.assertEqual(calculo.lista_campos_tabela(ind), ['Motivo'])
        self.assertEqual(calculo.resolver_aliases('Var 01 / Var 02', calculo.lista_campos(ind)), '[Perdidas] / [Produzidas]')

    def test_registro_por_colagem_distribui_numeros_pelas_variaveis(self):
        from kpi.models import ModeloColagem
        ind = Indicador.objects.create(nome='Quebras', meta=5, sentido='menor', exibir_semanal=True, exibir_diario=True,
                                       campos='Perdidas\nProduzidas', formula='[Perdidas] / [Produzidas] * 100')
        modelo = ModeloColagem.objects.create(nome='Fechamento', linhas='Produzidas | Quebras | Produzidas\nPedidos do dia\nPerdidas | Quebras | Perdidas')
        url = reverse('kpi:colagem')
        base = {'modelo': modelo.pk, 'granularidade': 'diaria', 'data': '2026-10-14', 'texto': '1123\n2139\n29'}
        r = self.client.post(url, base)  # conferir: não grava
        self.assertEqual(len(r.context['itens']), 1)
        self.assertEqual(str(r.context['itens'][0]['valor']), '2.58')
        self.assertFalse(Medicao.objects.filter(indicador=ind).exists())
        self.client.post(url, dict(base, confirmar='1'))
        dia = Medicao.objects.get(indicador=ind, granularidade='diaria')
        self.assertEqual(dia.dados, {'Produzidas': '1123', 'Perdidas': '29'})
        self.assertEqual(Medicao.objects.filter(indicador=ind, granularidade='semanal').count(), 1)  # consolidação
        # quantidade errada de números é recusada
        r = self.client.post(url, dict(base, texto='1\n2'))
        self.assertTrue(any('3 posições' in e for e in r.context['erros']))
        # modelo com variável inexistente é recusado ao salvar
        r = self.client.post(reverse('kpi:colagem_modelo_novo'), {'nome': 'Ruim', 'linhas': 'X | Quebras | Inexistente'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(ModeloColagem.objects.filter(nome='Ruim').exists())

    def test_modelos_e_importacao(self):
        import io
        from openpyxl import Workbook, load_workbook
        for tipo, aba in (('valores', 'Valores'), ('motivos', 'Motivos')):
            for gran in ('diaria', 'semanal', 'mensal'):
                r = self.client.get(reverse('kpi:modelo_importacao', args=[tipo, gran]))
                self.assertEqual(r.status_code, 200)
                wb = load_workbook(io.BytesIO(r.content))
                self.assertEqual(wb.sheetnames[0], aba)
                self.assertNotIn('Motivos' if tipo == 'valores' else 'Valores', wb.sheetnames)
        self.assertEqual(self.client.get(reverse('kpi:modelo_importacao', args=['valores', 'x'])).status_code, 404)
        self.assertEqual(self.client.get(reverse('kpi:modelo_importacao', args=['x', 'mensal'])).status_code, 404)
        self.assertEqual(self.client.get(reverse('kpi:importar')).status_code, 200)

        def enviar(wb, tipo, gran):
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            buf.name = 'x.xlsx'
            return self.client.post(reverse('kpi:importar'), {'tipo': tipo, 'granularidade': gran, 'arquivo': buf}).context['resultado']

        wb = Workbook()
        v = wb.active
        v.title = 'Valores'
        v.append(['Indicador', 'Semana', 'Valor', 'Observação'])
        v.append(['SR', '15/10/2026', '94,5', 'ok'])
        v.append(['Inexistente', '15/10/2026', 1, ''])
        res = enviar(wb, 'valores', 'semanal')
        self.assertEqual((res['valores'], res['motivos'], len(res['erros'])), (1, 0, 1))
        self.assertEqual(str(Medicao.objects.get(granularidade='semanal').valor), '94.50')
        # planilha de gráficos enviada como tabelas é recusada com mensagem clara
        res = enviar(wb, 'motivos', 'semanal')
        self.assertEqual((res['motivos'], len(res['erros'])), (0, 1))
        self.assertIn('Valores', res['erros'][0][2] + 'Valores')

        self.ind.motivos_semanal = True
        self.ind.save()
        wb = Workbook()
        m = wb.active
        m.title = 'Motivos'
        m.append(['Indicador', 'Semana', 'Motivo', 'Quantidade'])
        m.append(['SR', '13/10/2026', 'Quebra', 3])
        m.append(['SR', '14/10/2026', 'Trinca', 1])
        res = enviar(wb, 'motivos', 'semanal')
        self.assertEqual((res['valores'], res['motivos'], len(res['erros'])), (0, 2, 0))
        self.assertEqual(Motivo.objects.filter(granularidade='semanal', data=date(2026, 10, 12)).count(), 2)
        # e o contrário: tabelas enviadas como gráficos
        res = enviar(wb, 'valores', 'semanal')
        self.assertEqual((res['valores'], len(res['erros'])), (0, 1))
        self.assertEqual(Medicao.objects.filter(granularidade='semanal').count(), 1)

    def test_campos_formula_e_modelo_por_indicador(self):
        import io
        from openpyxl import Workbook, load_workbook
        ind = Indicador.objects.create(
            nome='Quebras', meta=5, sentido='menor', campos='Lentes produzidas\nLentes perdidas',
            formula='[Lentes perdidas] / [Lentes produzidas] * 100')
        # pop-up: calcula o valor pela fórmula quando vazio
        self.client.post(reverse('kpi:medicao_salvar', args=[ind.pk]), {
            'granularidade': 'mensal', 'periodo': '2026-10-10', 'valor': '', 'campo_0': '200', 'campo_1': '10', 'observacao': ''})
        m = Medicao.objects.get(indicador=ind)
        self.assertEqual(str(m.valor), '5.00')
        self.assertEqual(m.dados['Lentes perdidas'], '10')
        # fórmula inválida no cadastro
        r = self.client.post(reverse('kpi:indicador_editar', args=[ind.pk]), {
            'nome': 'Quebras', 'meta': '5', 'sentido': 'menor', 'periodicidade': 'mensal',
            'campos': 'A', 'formula': '[B] / 2'})
        self.assertEqual(r.status_code, 200)
        # modelo específico traz as colunas dos campos
        r = self.client.get(reverse('kpi:modelo_importacao', args=['valores', 'mensal']) + '?ind=%d' % ind.pk)
        wb = load_workbook(io.BytesIO(r.content))
        cab = [c.value for c in wb['Valores'][1]]
        self.assertEqual(cab[:4], ['Indicador', 'Mês', 'Lentes produzidas', 'Lentes perdidas'])
        # importação lê os campos pelo cabeçalho e calcula o valor
        wb = Workbook()
        v = wb.active
        v.title = 'Valores'
        v.append(['Indicador', 'Mês', 'Lentes produzidas', 'Lentes perdidas', 'Valor (opcional)', 'Observação'])
        v.append(['Quebras', '11/2026', 400, 8, None, ''])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'x.xlsx'
        r = self.client.post(reverse('kpi:importar'), {'tipo': 'valores', 'granularidade': 'mensal', 'arquivo': buf})
        self.assertEqual(len(r.context['resultado']['erros']), 0)
        novo = Medicao.objects.get(indicador=ind, periodo=date(2026, 11, 30))
        self.assertEqual(str(novo.valor), '2.00')
        self.assertEqual(novo.dados['Lentes produzidas'], '400')

    def test_flags_de_exibicao(self):
        base = {'nome': 'Flags', 'meta': '1', 'sentido': 'maior', 'periodicidade': 'mensal'}
        self.client.post(reverse('kpi:indicador_novo'), dict(base, ativo='on', exibir_mensal='on', exibir_diario='on', motivos_semanal='on'))
        ind = Indicador.objects.get(nome='Flags')
        self.assertEqual(ind.graficos(), ['mensal', 'diaria'])
        self.assertTrue(ind.permite_motivos('semanal') and not ind.permite_motivos('diaria'))
        r = self.client.get(reverse('kpi:painel') + '?ind=%d' % ind.pk)
        html = r.content.decode()
        pk = ind.pk
        self.assertIn('kpi-chart-%d-mensal' % pk, html)
        self.assertIn('kpi-chart-%d-diaria' % pk, html)
        self.assertNotIn('kpi-chart-%d-semanal' % pk, html)
        aba = [a for a in r.context['abas'] if a['ind'].pk == pk][0]
        self.assertEqual([t['gran'] for t in aba['motivos_tabs']], ['semanal'])
        # motivos de período não habilitado são recusados
        url = reverse('kpi:motivo_salvar', args=[pk])
        self.client.post(url, {'granularidade': 'diaria', 'data': '2026-10-15', 'descricao': ['X'], 'quantidade': ['1']})
        self.assertEqual(Motivo.objects.count(), 0)
        self.client.post(url, {'granularidade': 'semanal', 'data': '2026-10-15', 'descricao': ['X'], 'quantidade': ['1']})
        self.assertEqual(Motivo.objects.count(), 1)

    def test_setores_opcionais(self):
        import io
        from openpyxl import Workbook, load_workbook
        sem = Indicador.objects.create(nome='Sem setor', meta=5, sentido='menor', ativo=True)
        self.assertEqual(sem.setores.count(), 0)
        base = {'nome': 'Quebras', 'meta': '5', 'sentido': 'menor', 'periodicidade': 'mensal', 'ativo': 'on',
                'exibir_mensal': 'on', 'motivos_mensal': 'on'}
        r = self.client.post(reverse('kpi:indicador_novo'), dict(base, setores_texto='Surfaçagem | 2,5\nAR\n'))
        self.assertEqual(r.status_code, 302)
        ind = Indicador.objects.get(nome='Quebras')
        surf, ar = list(ind.setores.all())
        self.assertEqual((surf.nome, str(surf.meta), ar.meta), ('Surfaçagem', '2.50', None))
        # meta inválida / repetida
        r = self.client.post(reverse('kpi:indicador_editar', args=[ind.pk]), dict(base, setores_texto='A | x'))
        self.assertEqual(r.status_code, 200)
        # lançamento geral e por setor convivem no mesmo período
        url = reverse('kpi:medicao_salvar', args=[ind.pk])
        self.client.post(url, {'granularidade': 'mensal', 'periodo': '2026-10-05', 'valor': '4'})
        self.client.post(url, {'granularidade': 'mensal', 'periodo': '2026-10-05', 'valor': '3', 'setor': surf.pk})
        self.client.post(url, {'granularidade': 'mensal', 'periodo': '2026-10-09', 'valor': '2', 'setor': surf.pk})
        self.assertEqual(Medicao.objects.filter(indicador=ind).count(), 2)
        self.assertEqual(str(Medicao.objects.get(indicador=ind, setor=surf).valor), '2.00')
        # setor de outro indicador é recusado
        outro = Indicador.objects.create(nome='Outro', meta=1)
        self.client.post(reverse('kpi:medicao_salvar', args=[outro.pk]),
                         {'granularidade': 'mensal', 'periodo': '2026-10-05', 'valor': '1', 'setor': surf.pk})
        self.assertEqual(Medicao.objects.filter(indicador=outro).count(), 0)
        # painel: escopo geral x setor, com a meta do setor
        r = self.client.get(reverse('kpi:painel') + '?ind=%d' % ind.pk)
        aba = [a for a in r.context['abas'] if a['ind'].pk == ind.pk][0]
        self.assertEqual((aba['setor'], str(aba['ultima'].valor), str(aba['meta'])), (None, '4.00', '5.00'))
        r = self.client.get(reverse('kpi:painel') + '?ind=%d&setor=%d' % (ind.pk, surf.pk))
        aba = [a for a in r.context['abas'] if a['ind'].pk == ind.pk][0]
        self.assertEqual((aba['setor'], str(aba['ultima'].valor), str(aba['meta'])), (surf, '2.00', '2.50'))
        self.assertTrue(aba['atingiu_escopo'])
        # setor sem meta usa a geral
        r = self.client.get(reverse('kpi:painel') + '?ind=%d&setor=%d' % (ind.pk, ar.pk))
        aba = [a for a in r.context['abas'] if a['ind'].pk == ind.pk][0]
        self.assertEqual(str(aba['meta']), '5.00')
        # motivos por setor
        m_url = reverse('kpi:motivo_salvar', args=[ind.pk])
        self.client.post(m_url, {'granularidade': 'mensal', 'data': '2026-10-01', 'descricao': ['Geral'], 'quantidade': ['1']})
        self.client.post(m_url, {'granularidade': 'mensal', 'data': '2026-10-01', 'descricao': ['Do setor'], 'quantidade': ['2'], 'setor': surf.pk})
        self.assertEqual(Motivo.objects.filter(setor=surf).count(), 1)
        self.assertEqual(Motivo.objects.filter(setor__isnull=True).count(), 1)
        # visão geral soma os motivos de todos os setores; a do setor mostra só os dele
        def top(qs):
            aba = [a for a in self.client.get(reverse('kpi:painel') + qs).context['abas'] if a['ind'].pk == ind.pk][0]
            return {t['descricao']: float(t['total']) for t in aba['motivos_tabs'][0]['top']}
        self.assertEqual(top('?ind=%d&mot=0&p_mensal=acum' % ind.pk), {'Do setor': 2.0, 'Geral': 1.0})
        self.assertEqual(top('?ind=%d&setor=%d&mot=0&p_mensal=acum' % (ind.pk, surf.pk)), {'Do setor': 2.0})
        self.assertEqual(top('?ind=%d&setor=%d&mot=0&p_mensal=acum' % (ind.pk, ar.pk)), {})
        # filtro de período dos motivos: mês específico e intervalo
        def total(qs):
            aba = [a for a in self.client.get(reverse('kpi:painel') + qs).context['abas'] if a['ind'].pk == ind.pk][0]
            return sum(t['total'] for t in aba['motivos_tabs'][0]['top'])
        self.client.post(m_url, {'granularidade': 'mensal', 'data': '2026-09-15', 'descricao': ['Setembro'], 'quantidade': ['5']})
        self.assertEqual(total('?ind=%d&mot=mes&mes=2026-09&p_mensal=acum' % ind.pk), 5)
        self.assertEqual(total('?ind=%d&mot=mes&mes=2026-10&p_mensal=acum' % ind.pk), 3)
        self.assertEqual(total('?ind=%d&mot=periodo&de=2026-09-01&ate=2026-10-31&p_mensal=acum' % ind.pk), 8)
        self.assertEqual(total('?ind=%d&mot=periodo&de=2026-10-31&ate=2026-09-30&p_mensal=acum' % ind.pk), 8)  # invertido é corrigido
        self.assertEqual(total('?ind=%d&mot=periodo&de=2026-10-01&p_mensal=acum' % ind.pk), 3)
        self.assertEqual(total('?ind=%d&mot=mes&mes=lixo&p_mensal=acum' % ind.pk), 8)  # inválido cai em "todo o período"
        # setor com lançamentos não pode ser removido
        r = self.client.post(reverse('kpi:indicador_editar', args=[ind.pk]), dict(base, setores_texto='AR'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(ind.setores.count(), 2)
        # modelo traz a coluna Setor e a importação a respeita
        r = self.client.get(reverse('kpi:modelo_importacao', args=['valores', 'mensal']) + '?ind=%d' % ind.pk)
        self.assertIn('Setor', [c.value for c in load_workbook(io.BytesIO(r.content))['Valores'][1]])
        r = self.client.get(reverse('kpi:modelo_importacao', args=['motivos', 'mensal']) + '?ind=%d' % ind.pk)
        self.assertIn('Setor', [c.value for c in load_workbook(io.BytesIO(r.content))['Motivos'][1]])
        wb = Workbook()
        v = wb.active
        v.title = 'Valores'
        v.append(['Indicador', 'Mês', 'Setor', 'Valor', 'Observação'])
        v.append(['Quebras', '11/2026', 'AR', 1.5, ''])
        v.append(['Quebras', '11/2026', 'Inexistente', 1, ''])
        v.append(['Quebras', '11/2026', None, 4, ''])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'x.xlsx'
        r = self.client.post(reverse('kpi:importar'), {'tipo': 'valores', 'granularidade': 'mensal', 'arquivo': buf})
        self.assertEqual((r.context['resultado']['valores'], len(r.context['resultado']['erros'])), (2, 1))
        self.assertEqual(Medicao.objects.filter(indicador=ind, periodo=date(2026, 11, 30)).count(), 2)

    def test_navegacao_por_periodo_dos_motivos(self):
        ind = Indicador.objects.create(nome='Nav', meta=1, ativo=True, motivos_diario=True, motivos_semanal=True, motivos_mensal=True)
        url = reverse('kpi:motivo_salvar', args=[ind.pk])
        for dia, desc, q in (('2026-10-14', 'A', '4'), ('2026-10-15', 'B', '2'), ('2026-09-20', 'C', '7')):
            for gran in ('diaria', 'mensal'):
                self.client.post(url, {'granularidade': gran, 'data': dia, 'descricao': [desc], 'quantidade': [q]})
        self.client.post(url, {'granularidade': 'semanal', 'data': '2026-10-15', 'descricao': ['S'], 'quantidade': ['1']})

        def aba(qs=''):
            r = self.client.get(reverse('kpi:painel') + '?ind=%d%s' % (ind.pk, qs))
            return {t['gran']: t for t in [a for a in r.context['abas'] if a['ind'].pk == ind.pk][0]['motivos_tabs']}

        d = aba()
        # padrão: período mais recente com lançamentos em cada aba
        self.assertEqual([t['descricao'] for t in d['diaria']['top']], ['B'])
        self.assertIn('15/10/2026', d['diaria']['nav']['rotulo'])
        self.assertEqual([t['descricao'] for t in d['mensal']['top']], ['B'])  # B substituiu A (mesmo mês)
        self.assertIn('Semana 12/10', d['semanal']['nav']['rotulo'])
        # navegar para o dia anterior
        d = aba('&p_diaria=2026-10-14&mtab=diaria')
        self.assertEqual([t['descricao'] for t in d['diaria']['top']], ['A'])
        self.assertTrue(d['diaria']['ativo'] and not d['mensal']['ativo'])
        self.assertIn('p_diaria=2026-10-13', d['diaria']['nav']['prev'])
        self.assertIn('p_diaria=2026-10-15', d['diaria']['nav']['next'])
        self.assertIn('mtab=diaria', d['diaria']['nav']['prev'])
        # mês: anterior / seguinte atravessam o ano
        d = aba('&p_mensal=2026-01-31')
        self.assertIn('p_mensal=2025-12-31', d['mensal']['nav']['prev'])
        self.assertIn('p_mensal=2026-02-28', d['mensal']['nav']['next'])
        self.assertEqual(d['mensal']['top'], [])
        # semana é normalizada para segunda-feira e o rótulo mostra o intervalo
        d = aba('&p_semanal=2026-10-14')
        self.assertIn('12/10 – 18/10/2026', d['semanal']['nav']['rotulo'])
        # acumulado usa o filtro de período
        d = aba('&p_mensal=acum&mot=0')
        self.assertEqual(sorted(t['descricao'] for t in d['mensal']['top']), ['B', 'C'])
        self.assertTrue(d['mensal']['nav']['acum'])
        d = aba('&p_mensal=acum&mot=mes&mes=2026-09')
        self.assertEqual([t['descricao'] for t in d['mensal']['top']], ['C'])
        self.assertIn('09/2026', d['mensal']['nav']['rotulo'])

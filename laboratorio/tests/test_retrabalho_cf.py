import os
import tempfile
from datetime import date, datetime, time

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from laboratorio import cf_importacao as imp
from laboratorio import cf_retrabalho as cf
from laboratorio.models import (
    CFCadastro,
    CFCampoFormulario,
    EquipeCF,
    EquipeCoating,
    RegraTurnoCoating,
    RetrabalhoCFConfig,
    RetrabalhoCFRegistro,
    TratamentoAntiReflexo,
)
from maquinas.models import Maquina
from organization.models import Setor
from rh.models import Colaborador


def _colab(matricula, nome, ativo=True):
    setor, _ = Setor.objects.get_or_create(nome="Controle Final")
    return Colaborador.objects.create(matricula=matricula, nome_completo=nome, grupo="CF", setor=setor, is_active=ativo)


class CFBase(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("adm_cf", "a@a.com", "x")
        self.client.force_login(self.user)
        cf.garantir_padroes()
        self.t1 = RegraTurnoCoating.objects.create(nome="1º Turno", hora_inicio=time(6), hora_fim=time(14))
        self.t2 = RegraTurnoCoating.objects.create(nome="2º Turno", hora_inicio=time(14), hora_fim=time(22))
        self.gisele = _colab("1", "GISELE FREITAS")
        self.gecildo = _colab("2", "GECILDO DE MEDEIROS LIRA")
        self.montador = _colab("3", "MAILSON BASTOS")
        EquipeCF.objects.create(colaborador=self.gisele)
        EquipeCF.objects.create(colaborador=self.gecildo)
        EquipeCoating.objects.create(colaborador=self.montador, pode_montar=True)
        self.maq = Maquina.objects.create(codigo="DLX1200", nome="DLX 1200")
        self.trat = TratamentoAntiReflexo.objects.create(nome="Crizal Easy Pro")
        self.trat2 = TratamentoAntiReflexo.objects.create(nome="Crizal Rock")
        self.ar_risco = CFCadastro.objects.get(tipo="MOTIVO", grupo="AR", nome="Risco")
        self.hc_risco = CFCadastro.objects.get(tipo="MOTIVO", grupo="HC", nome="Risco")
        self.cf_fog = CFCadastro.objects.get(tipo="MOTIVO", grupo="CF", nome="FOG ID errado")

    def post_registro(self, **extra):
        dados = {"data": "2026-10-05", "turno": self.t1.pk, "numero_lote": "10203", "total_lentes": "100", "maquina": self.maq.pk,
                 "tratamento": self.trat.pk, "colaborador_cf": self.gisele.pk, "colaborador_ar": self.montador.pk,
                 "necessidade_retrabalho": "0", "lote_reprovado": "0"}
        dados.update({k: str(v) for k, v in extra.items()})
        return self.client.post(reverse("laboratorio:cf_registro_create"), dados)


class FormularioTests(CFBase):
    def test_lote_sem_retrabalho_guarda_nomes_das_integracoes(self):
        self.assertRedirects(self.post_registro(), reverse("laboratorio:cf_registros"))
        r = RetrabalhoCFRegistro.objects.get()
        self.assertEqual((r.total_lentes, r.qtd_retrabalho, r.reprovado), (100, 0, False))
        self.assertEqual((r.turno_nome, r.maquina_nome, r.tratamento_nome), ("1º Turno", "DLX 1200", "Crizal Easy Pro"))
        self.assertEqual(r.colaborador_cf_display, "GISELE FREITAS")
        self.assertEqual(r.colaborador_ar_display, "MAILSON BASTOS")

    def test_retrabalho_por_motivo_e_validacoes(self):
        self.post_registro(necessidade_retrabalho=1, **{f"mt_{self.ar_risco.pk}": 3, f"mt_{self.cf_fog.pk}": 2})
        r = RetrabalhoCFRegistro.objects.get()
        self.assertEqual((r.qtd_retrabalho, r.teve_retrabalho, r.itens.count()), (5, True, 2))
        # sem motivos / acima do total de lentes não gravam
        self.post_registro(necessidade_retrabalho=1)
        self.post_registro(necessidade_retrabalho=1, total_lentes=4, **{f"mt_{self.ar_risco.pk}": 9})
        self.assertEqual(RetrabalhoCFRegistro.objects.count(), 1)

    def test_reprovacao_so_vale_quando_lote_reprovado(self):
        motivo = CFCadastro.objects.get(tipo="REPROVACAO").nome
        self.post_registro(lote_reprovado=1, motivo_reprovacao=motivo)
        self.post_registro(lote_reprovado=0, motivo_reprovacao=motivo, numero_lote="2")
        reprovado = RetrabalhoCFRegistro.objects.get(numero_lote="10203")
        normal = RetrabalhoCFRegistro.objects.get(numero_lote="2")
        self.assertEqual((reprovado.reprovado, reprovado.motivo_reprovacao), (True, motivo))
        self.assertEqual((normal.reprovado, normal.motivo_reprovacao), (False, ""))

    def test_listas_vem_dos_cadastros_integrados(self):
        outro = _colab("9", "FORA DA LISTA")
        form = self.client.get(reverse("laboratorio:cf_registro_create")).context["form"]
        self.assertEqual([c.pk for c in form.fields["colaborador_cf"].queryset], [self.gecildo.pk, self.gisele.pk])
        self.assertEqual(list(form.fields["colaborador_ar"].queryset), [self.montador])
        self.assertNotIn(outro, form.fields["colaborador_ar"].queryset)
        self.assertEqual(list(form.fields["maquina"].queryset), [self.maq])

    def test_matriz_colaboradores_por_campo(self):
        from laboratorio.models import EquipeHC
        novo = _colab("20", "NOVO COLAB")
        url = reverse("laboratorio:cf_equipe_save")
        self.client.post(url, {"colaborador": novo.pk, "cf": "on", "hc": "on"})
        self.assertTrue(EquipeCF.objects.get(colaborador=novo).ativo)
        self.assertTrue(EquipeHC.objects.get(colaborador=novo).ativo)
        self.assertFalse(EquipeCoating.objects.filter(colaborador=novo, pode_montar=True).exists())
        # desmarcar CF e marcar montagem (sem dar acesso à Preparação)
        self.client.post(url, {"colaborador": novo.pk, "ar": "on", "hc": "on"})
        self.assertFalse(EquipeCF.objects.get(colaborador=novo).ativo)
        ar = EquipeCoating.objects.get(colaborador=novo)
        self.assertEqual((ar.pode_montar, ar.pode_preparar), (True, False))
        form = self.client.get(reverse("laboratorio:cf_registro_create")).context["form"]
        self.assertNotIn(novo, form.fields["colaborador_cf"].queryset)
        self.assertIn(novo, form.fields["colaborador_ar"].queryset)
        busca = self.client.get(reverse("laboratorio:cf_colaboradores_busca"), {"q": "novo"}).json()["results"]
        self.assertEqual(busca, [])  # já está na matriz

    def test_editar_e_excluir(self):
        self.post_registro(necessidade_retrabalho=1, **{f"mt_{self.ar_risco.pk}": 3})
        r = RetrabalhoCFRegistro.objects.get()
        resp = self.client.get(reverse("laboratorio:cf_registro_update", args=[r.pk]))
        self.assertEqual(resp.context["form"].initial[f"mt_{self.ar_risco.pk}"], 3)
        self.client.post(reverse("laboratorio:cf_registro_update", args=[r.pk]), {
            "data": "2026-10-05", "turno": self.t1.pk, "total_lentes": 50, "colaborador_cf": self.gisele.pk,
            "necessidade_retrabalho": "1", f"mt_{self.hc_risco.pk}": "2"})
        r.refresh_from_db()
        self.assertEqual((r.total_lentes, r.qtd_retrabalho), (50, 2))
        self.assertEqual(r.itens.get().motivo, self.hc_risco)
        self.client.post(reverse("laboratorio:cf_registro_delete", args=[r.pk]))
        self.assertFalse(RetrabalhoCFRegistro.objects.exists())


class DetalhePopupTests(CFBase):
    def test_detalhe_mostra_dados_e_acoes_so_no_popup(self):
        self.post_registro(necessidade_retrabalho=1, **{f"mt_{self.ar_risco.pk}": 3, f"mt_{self.cf_fog.pk}": 2})
        r = RetrabalhoCFRegistro.objects.get()
        html = self.client.get(reverse("laboratorio:cf_registro_detalhe", args=[r.pk])).content.decode()
        for trecho in ("10203", "GISELE FREITAS", "Anti-Reflexo", "FOG ID errado", reverse("laboratorio:cf_registro_update", args=[r.pk]),
                       reverse("laboratorio:cf_registro_delete", args=[r.pk])):
            self.assertIn(trecho, html)
        lista = self.client.get(reverse("laboratorio:cf_registros")).content.decode()
        self.assertIn(reverse("laboratorio:cf_registro_detalhe", args=[r.pk]), lista)
        self.assertNotIn(reverse("laboratorio:cf_registro_update", args=[r.pk]), lista)
        self.assertNotIn(reverse("laboratorio:cf_registro_delete", args=[r.pk]), lista)
        self.assertNotIn("<th>Origem</th>", lista)  # coluna removida
        self.assertIn("card-registro", lista)  # cards no mobile
        self.assertIn("DOMContentLoaded", lista)  # o bootstrap só existe depois do conteúdo


class ModeloEExportacaoTests(CFBase):
    def _ler(self, resp):
        import io

        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
        return wb, [list(r) for r in wb["Respostas"].iter_rows(values_only=True)]

    def test_modelo_tem_cabecalhos_e_instrucoes(self):
        wb, linhas = self._ler(self.client.get(reverse("laboratorio:cf_modelo")))
        self.assertEqual(len(linhas), 1)
        self.assertIn("Total de Lentes no Lote", linhas[0])
        self.assertIn("Retrabalho CF - CF - FOG ID errado", linhas[0])
        self.assertIn("Instruções", wb.sheetnames)

    def test_exporta_filtrado_e_reimporta(self):
        motivo = CFCadastro.objects.get(tipo="REPROVACAO").nome
        self.post_registro(necessidade_retrabalho=1, lote_reprovado=1, motivo_reprovacao=motivo, **{f"mt_{self.ar_risco.pk}": 3, f"mt_{self.cf_fog.pk}": 2})
        self.post_registro(data="2026-09-01", numero_lote="velho")
        wb, linhas = self._ler(self.client.get(reverse("laboratorio:cf_exportar")))
        self.assertEqual(len(linhas), 3)
        _wb, filtrado = self._ler(self.client.get(reverse("laboratorio:cf_exportar"), {"inicio": "2026-10-01"}))
        self.assertEqual(len(filtrado), 2)
        linha = dict(zip(filtrado[0], filtrado[1]))
        self.assertEqual((linha["Total de Lentes no Lote"], linha["Retrabalho AR - AR - Risco"], linha["Lote Reprovado?"], linha["Máquina"]), (100, 3, "Sim", "DLX 1200"))
        fd, caminho = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self.addCleanup(os.remove, caminho)
        wb.save(caminho)
        RetrabalhoCFRegistro.objects.all().delete()
        with open(caminho, "rb") as f:
            dados = imp.ler_planilha(f)
        resultado = imp.importar(dados["linhas"], imp.sugerir_mapa(dados["cabecalhos"]))
        self.assertEqual(resultado["importadas"], 2)
        r = RetrabalhoCFRegistro.objects.get(data=date(2026, 10, 5))
        self.assertEqual((r.qtd_retrabalho, r.reprovado, r.colaborador_cf, r.colaborador_ar, r.tratamento), (5, True, self.gisele, self.montador, self.trat))
        self.assertEqual(r.itens.count(), 2)


class ConfiguracaoTests(CFBase):
    def test_campo_personalizado_vai_para_extras(self):
        self.client.post(reverse("laboratorio:cf_campo_save"), {"rotulo": "Estufa", "tipo": "SELECAO", "opcoes": "E1\nE2", "obrigatorio": "on", "visivel": "on"})
        self.post_registro()
        self.assertFalse(RetrabalhoCFRegistro.objects.exists())
        self.post_registro(x_estufa="E2")
        self.assertEqual(RetrabalhoCFRegistro.objects.get().extras, {"estufa": "E2"})

    def test_protecoes_de_campos_de_sistema_e_meta(self):
        qtd = CFCampoFormulario.objects.get(chave="total_lentes")
        self.client.post(reverse("laboratorio:cf_campo_save"), {"id": qtd.pk, "rotulo": "Lentes do lote"})
        qtd.refresh_from_db()
        self.assertTrue(qtd.visivel)
        self.client.post(reverse("laboratorio:cf_campo_delete", args=[qtd.pk]))
        self.assertTrue(CFCampoFormulario.objects.filter(pk=qtd.pk).exists())
        self.client.post(reverse("laboratorio:cf_config_meta"), {"meta_retrabalho": "3,5"})
        self.assertEqual(float(RetrabalhoCFConfig.get().meta_retrabalho), 3.5)
        self.client.post(reverse("laboratorio:cf_config_meta"), {"meta_retrabalho": "0"})
        self.assertEqual(float(RetrabalhoCFConfig.get().meta_retrabalho), 3.5)

    def test_novo_motivo_aparece_no_formulario_e_em_uso_so_inativa(self):
        self.client.post(reverse("laboratorio:cf_cadastro_save"), {"tipo": "MOTIVO", "grupo": "CF", "nome": "Lente trocada"})
        novo = CFCadastro.objects.get(nome="Lente trocada")
        self.assertEqual(novo.grupo, "CF")
        self.assertIn(f"mt_{novo.pk}", self.client.get(reverse("laboratorio:cf_registro_create")).context["form"].fields)
        self.post_registro(necessidade_retrabalho=1, **{f"mt_{novo.pk}": 1})
        self.client.post(reverse("laboratorio:cf_cadastro_delete", args=[novo.pk]))
        novo.refresh_from_db()
        self.assertFalse(novo.ativo)

    def test_paginas_abrem(self):
        for nome in ("cf_config_campos", "cf_config_cadastros", "cf_config_integracoes", "cf_registros", "cf_registro_create", "cf_importar"):
            self.assertEqual(self.client.get(reverse(f"laboratorio:{nome}")).status_code, 200, nome)
        for aba in ("MOTIVO", "REPROVACAO"):  # a aba EQUIPE agora redireciona para a Equipe HMC
            self.assertEqual(self.client.get(reverse("laboratorio:cf_config_cadastros"), {"aba": aba}).status_code, 200)


class DashboardTests(CFBase):
    def setUp(self):
        super().setUp()
        motivo = CFCadastro.objects.get(tipo="REPROVACAO").nome
        # dia 5: lote A (T1, Gisele, Easy Pro) 100 lentes, 5 retrabalho; lote B (T2, Gecildo, Rock) 100 lentes reprovado
        self.post_registro(necessidade_retrabalho=1, **{f"mt_{self.ar_risco.pk}": 3, f"mt_{self.cf_fog.pk}": 2})
        self.post_registro(numero_lote="2", turno=self.t2.pk, colaborador_cf=self.gecildo.pk, tratamento=self.trat2.pk,
                           lote_reprovado=1, motivo_reprovacao=motivo)

    def test_geral(self):
        d = cf.geral(date(2026, 10, 1), date(2026, 10, 31))
        dia5 = next(l for l in d["linhas"] if l["data"] == date(2026, 10, 5))
        self.assertEqual((dia5["lentes"], dia5["lotes"], dia5["ret"], dia5["pct_ret"]), (200, 2, 5, 2.5))
        self.assertEqual((dia5["lentes_rep"], dia5["pct_lentes_rep"]), (100, 50.0))
        self.assertEqual((dia5["ar"], dia5["hc"], dia5["cf"]), (3, 0, 2))
        self.assertIsNone(next(l for l in d["linhas"] if l["data"] == date(2026, 10, 6))["pct_ret"])

    def test_motivo_dimensoes_e_colaboradores(self):
        m = cf.retrabalho_motivo(date(2026, 10, 1), date(2026, 10, 31))
        ar = next(g for g in m["grupos"] if g["origem"] == "AR")
        self.assertEqual(next(l for l in ar["linhas"] if l["nome"] == "Risco")["turnos"][0]["qtd"], 3)
        self.assertEqual(m["total"], 5)
        maq = cf.analise_dimensao("maquina", "DLX 1200", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(maq["total"]["lentes"], 200)
        trat = cf.analise_dimensao("tratamento", "", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual({r["nome"]: r["pct_ret"] for r in trat["resumo"]}, {"Crizal Easy Pro": 5.0, "Crizal Rock": 0.0})
        col = cf.analise_colaborador("cf", f"id:{self.gecildo.pk}", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual((col["total"]["lotes_rep"], col["total"]["pct_lotes_rep"]), (1, 100.0))
        ar = cf.analise_colaborador("ar", "", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(ar["resumo"][0]["nome"], "MAILSON BASTOS")
        cxt = cf.colaborador_x_tratamento("cf", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(cxt["tratamentos"], ["Crizal Easy Pro", "Crizal Rock"])

    def test_todas_as_abas_renderizam(self):
        url = reverse("laboratorio:cf_dashboard")
        extras = {"colaborador": {"colaborador": f"id:{self.gisele.pk}"}, "maquina": {"item": "DLX 1200"}, "tratamento": {"item": "Crizal Rock"}}
        for aba in ("geral", "motivo", "maquina", "tratamento", "colaborador", "cxt"):
            self.assertEqual(self.client.get(url, {"aba": aba, **extras.get(aba, {})}).status_code, 200, aba)
        self.assertEqual(self.client.get(url, {"aba": "colaborador", "papel": "ar"}).status_code, 200)
        self.assertContains(self.client.get(url, {"aba": "geral"}), "2,50%")

    def test_dashboard_vazio(self):
        RetrabalhoCFRegistro.objects.all().delete()
        for aba in ("geral", "motivo", "maquina", "colaborador", "cxt"):
            self.assertEqual(self.client.get(reverse("laboratorio:cf_dashboard"), {"aba": aba, "colaborador": "lixo"}).status_code, 200)


class IntegracaoTests(CFBase):
    def test_sincronizacao_e_vinculos(self):
        novo = _colab("10", "NOVO INSPETOR")
        RetrabalhoCFRegistro.objects.create(data=date(2026, 10, 1), total_lentes=10, colaborador_cf=novo, colaborador_cf_nome=novo.nome_completo)
        inativo = _colab("11", "SAIU", ativo=False)
        EquipeCF.objects.create(colaborador=inativo)
        p = cf.propostas_sincronizacao()
        self.assertEqual(p["adicionar"], [novo])
        self.assertEqual([e.colaborador for e in p["inativos"]], [inativo])
        self.client.post(reverse("laboratorio:cf_sincronizar"))
        self.assertTrue(EquipeCF.objects.get(colaborador=novo).ativo)
        self.assertFalse(EquipeCF.objects.get(colaborador=inativo).ativo)

    def test_vincular_nomes_antigos(self):
        RetrabalhoCFRegistro.objects.create(data=date(2026, 10, 1), total_lentes=10, colaborador_cf_nome="Gise", maquina_nome="380 A", tratamento_nome="No Reflex")
        self.client.post(reverse("laboratorio:cf_vincular"), {"chave": "colab_cf", "nome": "Gise", "objeto": self.gisele.pk})
        self.client.post(reverse("laboratorio:cf_vincular"), {"chave": "maquina", "nome": "380 A", "objeto": self.maq.pk})
        self.client.post(reverse("laboratorio:cf_criar_tratamento"), {"nome": "No Reflex"})
        r = RetrabalhoCFRegistro.objects.get()
        self.assertEqual((r.colaborador_cf, r.maquina, r.maquina_nome, r.tratamento.nome), (self.gisele, self.maq, "DLX 1200", "No Reflex"))
        self.assertIn("Gise", EquipeCF.objects.get(colaborador=self.gisele).apelidos)
        self.assertFalse(any(v["itens"] for v in cf.propostas_sincronizacao()["vinculos"].values()))


class ImportacaoTests(CFBase):
    def _planilha(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Respostas"
        ws.append(["Id", "Start time", "Completion time", "Email", "Name", "Turno", "Data", "Número do Lote", "Total de Lentes no Lote", "Máquina",
                   "Tratamento", "Colaborador CF - Responsável Inspeção do Lote", "Necessidade Retrabalho?",
                   "Retrabalho AR\xa0 - AR - Batida _ Risco\xa0 de Anel", "​Retrabalho AR\xa0 - AR - Risco",
                   "Retrabalho HC\xa0-\xa0HC - Risco", "​Retrabalho HC\xa0-\xa0HC - Outros motivos HC",
                   "Retrabalho CF\xa0-\xa0CF - FOG ID errado", "Lote Reprovado?", "Motivo da Reprovação",
                   "Colaborador AR - Responsável Montagem Lote"])
        ws.append([1, datetime(2026, 10, 5, 9), None, "x", None, "1º Turno", date(2026, 10, 5), 10203, 70, "DLX 1200", "Crizal Easy Pro", "Gisele", "Sim",
                   None, 2, 1, None, 1, "Não", None, "Mailson"])
        ws.append([2, datetime(2026, 10, 8, 9), None, "x", None, "2º Turno", date(2026, 8, 10), 10204, 60, "DLX 1200", "Crizal Rock", "Gecildo", "Não",
                   None, None, None, None, None, "Sim", "Lote Reprovado - Outros Motivos", "Mailson"])  # dia/mês invertidos
        ws.append([3, datetime(2026, 10, 8, 9), None, "x", None, "2º Turno", date(2001, 3, 1), 1, 5])  # data inválida
        fd, caminho = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(caminho)
        self.addCleanup(os.remove, caminho)
        return caminho

    def _upload(self, url, caminho, **extra):
        with open(caminho, "rb") as f:
            return self.client.post(reverse(url), {"arquivo": f, **extra})

    def test_sugestao_mapeia_motivos_pela_origem(self):
        d = self._upload("laboratorio:cf_importar_preview", self._planilha()).json()
        s, cab = d["sugestao"], d["cabecalhos"]
        self.assertEqual(cab[s["total_lentes"]], "Total de Lentes no Lote")
        self.assertEqual(cab[s["data_ref"]], "Start time")
        self.assertEqual(cab[s[f"mt_{self.ar_risco.pk}"]].strip("​").split(" - ")[-1], "Risco")
        self.assertIn("Batida", cab[s[f"mt_{CFCadastro.objects.get(grupo='AR', nome='Batida / Risco de Anel').pk}"]])
        self.assertIn("HC", cab[s[f"mt_{self.hc_risco.pk}"]])
        self.assertIn("FOG ID errado", cab[s[f"mt_{self.cf_fog.pk}"]])
        self.assertEqual(cab[s["colaborador_ar"]], "Colaborador AR - Responsável Montagem Lote")

    def test_importa_corrige_data_e_vincula_cadastros(self):
        caminho = self._planilha()
        mapa = {f"map_{k}": v for k, v in self._upload("laboratorio:cf_importar_preview", caminho).json()["sugestao"].items()}
        resp = self._upload("laboratorio:cf_importar", caminho, acao="simular", **mapa)
        self.assertEqual((resp.context["resultado"]["validas"], resp.context["resultado"]["datas_corrigidas"]), (2, 1))
        self.assertFalse(RetrabalhoCFRegistro.objects.exists())
        resp = self._upload("laboratorio:cf_importar", caminho, acao="importar", **mapa)
        self.assertEqual(resp.context["resultado"]["importadas"], 2)
        a = RetrabalhoCFRegistro.objects.get(chave_importacao="id:1")
        self.assertEqual((a.qtd_retrabalho, a.colaborador_cf, a.colaborador_ar, a.maquina_nome, a.tratamento), (4, self.gisele, self.montador, "DLX 1200", self.trat))
        self.assertEqual(a.turno, self.t1)
        b = RetrabalhoCFRegistro.objects.get(chave_importacao="id:2")
        self.assertEqual((b.data, b.reprovado, b.motivo_reprovacao), (date(2026, 10, 8), True, "Lote Reprovado - Outros Motivos"))
        # reimportar não duplica
        self._upload("laboratorio:cf_importar", caminho, acao="importar", **mapa)
        self.assertEqual(RetrabalhoCFRegistro.objects.count(), 2)

    def test_sem_coluna_obrigatoria_nao_importa(self):
        caminho = self._planilha()
        mapa = {f"map_{k}": v for k, v in self._upload("laboratorio:cf_importar_preview", caminho).json()["sugestao"].items() if k != "total_lentes"}
        resp = self._upload("laboratorio:cf_importar", caminho, acao="importar", **mapa)
        self.assertIsNone(resp.context["resultado"])
        self.assertFalse(RetrabalhoCFRegistro.objects.exists())


class PermissoesCFTests(CFBase):
    def test_views_cf_no_mapa_de_permissoes(self):
        from shared.permissions import get_view_permission_map

        mapa = get_view_permission_map()
        for nome in ("cf_dashboard", "cf_registros", "cf_registro_create", "cf_registro_update", "cf_registro_delete", "cf_importar",
                     "cf_importar_preview", "cf_config_campos", "cf_campo_save", "cf_campo_mover", "cf_campo_delete", "cf_config_meta",
                     "cf_config_cadastros", "cf_cadastro_save", "cf_cadastro_delete", "cf_equipe_save", "cf_equipe_delete",
                     "cf_config_integracoes", "cf_sincronizar", "cf_vincular", "cf_criar_tratamento", "cf_modelo", "cf_exportar"):
            self.assertIn(f"laboratorio:{nome}", mapa, nome)

    def test_menu_mostra_itens_do_cf(self):
        resp = self.client.get(reverse("laboratorio:cf_dashboard"))
        self.assertContains(resp, reverse("laboratorio:cf_config_campos"))
        self.assertContains(resp, reverse("laboratorio:cf_registro_create"))


class DiaHoraCFTests(CFBase):
    def _reg(self, hora, total, ret=0, **kw):
        from django.utils import timezone

        r = RetrabalhoCFRegistro.objects.create(data=date(2026, 10, 5), total_lentes=total, qtd_retrabalho=ret, teve_retrabalho=bool(ret), **kw)
        if hora is not None:
            RetrabalhoCFRegistro.objects.filter(pk=r.pk).update(
                criado_em=timezone.make_aware(datetime(2026, 10, 5, hora, 30)))
        return r

    def test_agrupa_por_hora_e_isola_registros_sem_hora(self):
        from laboratorio.models import RetrabalhoCFItem

        self._reg(8, 100, 4)
        r = self._reg(8, 100, 0, reprovado=True)
        self._reg(10, 50, 5)
        RetrabalhoCFItem.objects.create(registro=RetrabalhoCFRegistro.objects.get(total_lentes=50), motivo=self.ar_risco, quantidade=5)
        self._reg(None, 200, 2, origem=RetrabalhoCFRegistro.ORIGEM_IMPORT)
        d = cf.por_hora(date(2026, 10, 5))
        self.assertEqual([l["rotulo"] for l in d["linhas"]], ["08h", "09h", "10h"])
        h8, h9, h10 = d["linhas"]
        self.assertEqual((h8["lentes"], h8["lotes"], h8["ret"], h8["pct_ret"], h8["lentes_rep"]), (200, 2, 4, 2.0, 100))
        self.assertEqual((h9["lotes"], h9["pct_ret"]), (0, None))
        self.assertEqual((h10["ar"], h10["pct_ar"], h10["pct_ret"]), (5, 10.0, 10.0))
        self.assertEqual((d["sem_hora"]["lentes"], d["total"]["lentes"], d["total"]["ret"]), (200, 450, 11))

    def test_tela_renderiza(self):
        url = reverse("laboratorio:cf_dia_hora")
        self.assertEqual(self.client.get(url).status_code, 200)
        self._reg(14, 100, 5)
        resp = self.client.get(url, {"data": "2026-10-05"})
        self.assertContains(resp, "14h")
        self.assertContains(resp, "5,00%")
        self.assertEqual(self.client.get(url, {"data": "lixo"}).status_code, 200)

    def test_view_no_mapa_de_permissoes(self):
        from shared.permissions import get_view_permission_map

        self.assertIn("laboratorio:cf_dia_hora", get_view_permission_map())


class HoraRegistroCFTests(CFBase):
    def test_formulario_grava_data_e_hora(self):
        resp = self.post_registro(data="2026-10-05T14:35")
        self.assertEqual(resp.status_code, 302)
        r = RetrabalhoCFRegistro.objects.get()
        self.assertEqual((r.data, r.hora), (date(2026, 10, 5), time(14, 35)))
        d = cf.por_hora(date(2026, 10, 5))
        self.assertEqual([l["rotulo"] for l in d["linhas"]], ["14h"])

    def test_edicao_mantem_hora_e_data_sem_hora_fica_sem_hora(self):
        self.post_registro(data="2026-10-05")
        r = RetrabalhoCFRegistro.objects.get()
        self.assertIsNone(r.hora)
        self.assertContains(self.client.get(reverse("laboratorio:cf_registro_update", args=[r.pk])), 'type="datetime-local"')

    def test_hora_na_importacao_e_reimportacao_preenche_existentes(self):
        mapa = {"data": 0, "total_lentes": 1, "data_ref": 2, "id": 3}
        linhas = [[date(2026, 10, 5), 100, datetime(2026, 10, 5, 9, 41), 77]]
        imp.importar(linhas, {**mapa, "data_ref": 9})
        self.assertIsNone(RetrabalhoCFRegistro.objects.get().hora)
        res = imp.importar(linhas, mapa)
        self.assertEqual(res["horas_preenchidas"], 1)
        self.assertEqual(RetrabalhoCFRegistro.objects.get().hora, time(9, 41))


class EquipeHMCUnificadaTests(CFBase):
    def test_tela_unica_com_colunas_dos_quatro_modulos(self):
        from laboratorio.models import EquipeHC
        EquipeHC.objects.create(colaborador=self.montador, ativo=True)
        r = self.client.get(reverse("laboratorio:equipe_coating_list"))
        self.assertEqual(r.status_code, 200)
        linhas = {l["colaborador"].pk: l for l in r.context["equipe"]}
        self.assertEqual(set(linhas), {self.gisele.pk, self.gecildo.pk, self.montador.pk})
        self.assertTrue(linhas[self.gisele.pk]["cf"] and not linhas[self.gisele.pk]["montar"] and not linhas[self.gisele.pk]["hc"])
        m = linhas[self.montador.pk]
        self.assertTrue(m["montar"] and m["hc"] and not m["cf"])
        # as abas "equipe" antigas levam à tela única
        for nome in ("cf_config_cadastros", "hc_config_cadastros"):
            resp = self.client.get(reverse(f"laboratorio:{nome}") + "?aba=EQUIPE")
            self.assertRedirects(resp, reverse("laboratorio:equipe_coating_list"), fetch_redirect_response=False)

    def test_adicionar_editar_e_remover_gravam_nos_tres_cadastros(self):
        from laboratorio.models import EquipeHC
        novo = _colab("9", "NOVO COLAB")
        r = self.client.post(reverse("laboratorio:equipe_coating_list"), {
            "colaborador": novo.pk, "preparar": "on", "cf": "on", "apelidos": "Novinho"})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(EquipeCoating.objects.get(colaborador=novo).pode_preparar)
        self.assertFalse(EquipeCoating.objects.get(colaborador=novo).pode_montar)
        self.assertEqual(EquipeCF.objects.get(colaborador=novo).apelidos, "Novinho")
        self.assertFalse(EquipeHC.objects.filter(colaborador=novo).exists())  # coluna desmarcada de quem nunca esteve nela
        # editar: troca as colunas; desmarcar mantém o cadastro como inativo
        self.client.post(reverse("laboratorio:equipe_coating_update", args=[novo.pk]), {"montar": "on", "hc": "on", "apelidos": "N2"})
        eq = EquipeCoating.objects.get(colaborador=novo)
        self.assertEqual((eq.pode_preparar, eq.pode_montar), (False, True))
        self.assertFalse(EquipeCF.objects.get(colaborador=novo).ativo)
        self.assertTrue(EquipeHC.objects.get(colaborador=novo).ativo)
        self.assertEqual(EquipeCF.objects.get(colaborador=novo).apelidos, "N2")
        # a busca Ajax já não oferece quem está na equipe; remover limpa tudo
        busca = self.client.get(reverse("laboratorio:cf_colaboradores_busca"), {"q": "NOVO"}).json()["results"]
        self.assertEqual(busca, [])
        self.client.post(reverse("laboratorio:equipe_coating_delete", args=[novo.pk]))
        self.assertFalse(EquipeCoating.objects.filter(colaborador=novo).exists())
        self.assertFalse(EquipeCF.objects.filter(colaborador=novo).exists())
        self.assertFalse(EquipeHC.objects.filter(colaborador=novo).exists())
        busca = self.client.get(reverse("laboratorio:cf_colaboradores_busca"), {"q": "NOVO"}).json()["results"]
        self.assertEqual(len(busca), 1)

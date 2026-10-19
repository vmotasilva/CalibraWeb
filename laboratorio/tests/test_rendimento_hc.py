import os
import tempfile
from datetime import date, time

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from laboratorio import hc_rendimento as hc
from laboratorio.models import (
    EquipeCoating,
    EquipeHC,
    HCCadastro,
    HCCampoFormulario,
    RegraTurnoCoating,
    RendimentoHCConfig,
    RendimentoHCRegistro,
)
from organization.models import Setor
from rh.models import Colaborador


def _colab(matricula, nome, ativo=True):
    setor, _ = Setor.objects.get_or_create(nome="Laboratorio")
    return Colaborador.objects.create(matricula=matricula, nome_completo=nome, grupo="Lab", setor=setor, is_active=ativo)


class RendimentoHCBase(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("adm_hc", "a@a.com", "x")
        self.client.force_login(self.user)
        hc.garantir_padroes()
        self.turno1 = RegraTurnoCoating.objects.create(nome="1º Turno", hora_inicio=time(6), hora_fim=time(14))
        self.turno2 = RegraTurnoCoating.objects.create(nome="2º Turno", hora_inicio=time(14), hora_fim=time(22))
        self.ana = _colab("1", "ANA PAULA SOUZA")
        self.bia = _colab("2", "BIA LIMA")
        EquipeHC.objects.create(colaborador=self.ana)
        EquipeHC.objects.create(colaborador=self.bia)
        self.nova = HCCadastro.objects.get(tipo="GARRA", nome="Lente Nova")
        self.retr1 = HCCadastro.objects.get(tipo="GARRA", nome="1º Retrabalho")
        self.orma = HCCadastro.objects.get(tipo="MATERIAL", nome__startswith="Orma")
        self.risco = HCCadastro.objects.get(tipo="MOTIVO", nome="Risco")
        self.verniz = HCCadastro.objects.get(tipo="MOTIVO", nome="Ponto de Verniz")

    def post_registro(self, **extra):
        dados = {
            "data": "2026-10-05", "turno": self.turno1.pk, "colaborador": self.ana.pk, "numero_garra": "7",
            "tipo_garra": self.nova.pk, "qtd_lentes": "20", "teve_retrabalho": "0",
        }
        dados.update({k: str(v) for k, v in extra.items()})
        return self.client.post(reverse("laboratorio:hc_registro_create"), dados)


class FormularioTests(RendimentoHCBase):
    def test_lancamento_sem_retrabalho(self):
        resp = self.post_registro(**{f"pm_{self.orma.pk}": 20})
        self.assertRedirects(resp, reverse("laboratorio:hc_registros"))
        r = RendimentoHCRegistro.objects.get()
        self.assertEqual((r.qtd_lentes, r.qtd_retrabalho, r.teve_retrabalho), (20, 0, False))
        self.assertEqual(r.turno_nome, "1º Turno")
        self.assertEqual(r.itens.get().quantidade, 20)

    def test_lancamento_com_retrabalho_soma_motivos(self):
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2, f"rt_{self.verniz.pk}": 1, f"rm_{self.orma.pk}": 3})
        r = RendimentoHCRegistro.objects.get()
        self.assertEqual(r.qtd_retrabalho, 3)
        self.assertTrue(r.teve_retrabalho)
        self.assertEqual(r.itens.filter(tipo="RT").count(), 2)

    def test_validacoes_de_retrabalho(self):
        # sem quantidade informada
        self.post_registro(teve_retrabalho=1)
        # maior que a produção
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 25})
        # soma por motivo difere da soma por material
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2, f"rm_{self.orma.pk}": 5})
        self.assertEqual(RendimentoHCRegistro.objects.count(), 0)

    def test_retrabalho_zerado_quando_nao_teve(self):
        self.post_registro(teve_retrabalho=0, **{f"rt_{self.risco.pk}": 4})
        r = RendimentoHCRegistro.objects.get()
        self.assertEqual(r.qtd_retrabalho, 0)
        self.assertFalse(r.itens.filter(tipo="RT").exists())

    def test_editar_e_excluir(self):
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2})
        r = RendimentoHCRegistro.objects.get()
        resp = self.client.get(reverse("laboratorio:hc_registro_update", args=[r.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["form"].initial[f"rt_{self.risco.pk}"], 2)
        dados = {"data": "2026-10-05", "turno": self.turno1.pk, "colaborador": self.ana.pk, "tipo_garra": self.nova.pk,
                 "qtd_lentes": "30", "teve_retrabalho": "1", f"rt_{self.verniz.pk}": "5"}
        self.client.post(reverse("laboratorio:hc_registro_update", args=[r.pk]), dados)
        r.refresh_from_db()
        self.assertEqual((r.qtd_lentes, r.qtd_retrabalho), (30, 5))
        self.assertEqual(r.itens.filter(tipo="RT").count(), 1)
        self.client.post(reverse("laboratorio:hc_registro_delete", args=[r.pk]))
        self.assertEqual(RendimentoHCRegistro.objects.count(), 0)

    def test_colaborador_do_formulario_vem_da_equipe_hc(self):
        _colab("3", "FORA DA EQUIPE")
        resp = self.client.get(reverse("laboratorio:hc_registro_create"))
        nomes = [c.nome_completo for c in resp.context["form"].fields["colaborador"].queryset]
        self.assertEqual(nomes, ["ANA PAULA SOUZA", "BIA LIMA"])


class DetalhePopupTests(RendimentoHCBase):
    def test_detalhe_mostra_dados_e_acoes_so_no_popup(self):
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2, f"rm_{self.orma.pk}": 2, f"pm_{self.orma.pk}": 20})
        r = RendimentoHCRegistro.objects.get()
        html = self.client.get(reverse("laboratorio:hc_registro_detalhe", args=[r.pk])).content.decode()
        for trecho in ("ANA PAULA SOUZA", "Risco", "Retrabalho por Motivo", reverse("laboratorio:hc_registro_update", args=[r.pk]),
                       reverse("laboratorio:hc_registro_delete", args=[r.pk])):
            self.assertIn(trecho, html)
        lista = self.client.get(reverse("laboratorio:hc_registros")).content.decode()
        self.assertIn(reverse("laboratorio:hc_registro_detalhe", args=[r.pk]), lista)
        self.assertNotIn(reverse("laboratorio:hc_registro_update", args=[r.pk]), lista)
        self.assertNotIn(reverse("laboratorio:hc_registro_delete", args=[r.pk]), lista)
        self.assertNotIn("<th>Origem</th>", lista)  # coluna removida
        self.assertIn("card-registro", lista)  # cards no mobile
        self.assertIn("DOMContentLoaded", lista)  # o bootstrap só existe depois do conteúdo
        self.assertEqual(self.client.get(reverse("laboratorio:hc_registro_detalhe", args=[999999])).status_code, 404)


class ModeloEExportacaoTests(RendimentoHCBase):
    def _ler(self, resp):
        import io

        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
        return wb, [list(r) for r in wb["Respostas"].iter_rows(values_only=True)]

    def test_modelo_tem_cabecalhos_e_instrucoes(self):
        resp = self.client.get(reverse("laboratorio:hc_modelo"))
        self.assertIn("spreadsheetml", resp["Content-Type"])
        wb, linhas = self._ler(resp)
        self.assertEqual(len(linhas), 1)  # só o cabeçalho
        self.assertIn("Quantidade de Lentes", linhas[0])
        self.assertIn("Retrabalho Risco", linhas[0])
        self.assertIn("Instruções", wb.sheetnames)

    def test_exporta_filtrado_e_reimporta(self):
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2, f"rm_{self.orma.pk}": 2, f"pm_{self.orma.pk}": 20})
        self.post_registro(data="2026-09-01", numero_lote="x")
        wb, linhas = self._ler(self.client.get(reverse("laboratorio:hc_exportar")))
        self.assertEqual(len(linhas), 3)
        _wb, filtrado = self._ler(self.client.get(reverse("laboratorio:hc_exportar"), {"inicio": "2026-10-01"}))
        self.assertEqual(len(filtrado), 2)
        cab = filtrado[0]
        linha = dict(zip(cab, filtrado[1]))
        self.assertEqual((linha["Colaborador"], linha["Quantidade de Lentes"], linha["Retrabalho Risco"], linha["Teve Retrabalho ?"]), ("ANA PAULA SOUZA", 20, 2, "Sim"))
        # round-trip: apaga tudo e importa o arquivo exportado com o mapeamento sugerido
        fd, caminho = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self.addCleanup(os.remove, caminho)
        wb.save(caminho)
        RendimentoHCRegistro.objects.all().delete()
        call_command("importar_rendimento_hc", caminho, verbosity=0)
        r = RendimentoHCRegistro.objects.get(data=date(2026, 10, 5))
        self.assertEqual((r.qtd_lentes, r.qtd_retrabalho, r.colaborador, r.tipo_garra), (20, 2, self.ana, self.nova))
        self.assertEqual(RendimentoHCRegistro.objects.count(), 2)


class CamposConfiguraveisTests(RendimentoHCBase):
    def test_campo_personalizado_aparece_e_grava_em_extras(self):
        self.client.post(reverse("laboratorio:hc_campo_save"), {
            "rotulo": "Lote do verniz", "tipo": "SELECAO", "opcoes": "A1\nB2", "obrigatorio": "on", "visivel": "on"})
        campo = HCCampoFormulario.objects.get(chave="lote-do-verniz")
        self.assertFalse(campo.sistema)
        # obrigatório: sem resposta não grava
        self.post_registro()
        self.assertEqual(RendimentoHCRegistro.objects.count(), 0)
        self.post_registro(**{"x_lote-do-verniz": "B2"})
        self.assertEqual(RendimentoHCRegistro.objects.get().extras, {"lote-do-verniz": "B2"})

    def test_ocultar_campo_e_protecoes_de_sistema(self):
        trat = HCCampoFormulario.objects.get(chave="tratamento")
        self.client.post(reverse("laboratorio:hc_campo_save"), {"id": trat.pk, "rotulo": "Tratamento", "visivel": "on"})
        resp = self.client.get(reverse("laboratorio:hc_registro_create"))
        self.assertIn("tratamento", resp.context["form"].fields)
        # campo essencial não pode ser ocultado nem excluído
        qtd = HCCampoFormulario.objects.get(chave="qtd_lentes")
        self.client.post(reverse("laboratorio:hc_campo_save"), {"id": qtd.pk, "rotulo": "Qtd. de lentes"})
        qtd.refresh_from_db()
        self.assertTrue(qtd.visivel)
        self.assertEqual(qtd.rotulo, "Qtd. de lentes")
        self.client.post(reverse("laboratorio:hc_campo_delete", args=[qtd.pk]))
        self.assertTrue(HCCampoFormulario.objects.filter(pk=qtd.pk).exists())

    def test_reordenar_campos(self):
        primeiro, segundo = list(HCCampoFormulario.objects.all()[:2])
        self.client.post(reverse("laboratorio:hc_campo_mover", args=[primeiro.pk]), {"direcao": "down"})
        self.assertEqual(list(HCCampoFormulario.objects.all()[:2]), [segundo, primeiro])

    def test_meta(self):
        self.client.post(reverse("laboratorio:hc_config_meta"), {"meta_rendimento": "96,5"})
        self.assertEqual(float(RendimentoHCConfig.get().meta_rendimento), 96.5)
        self.client.post(reverse("laboratorio:hc_config_meta"), {"meta_rendimento": "150"})
        self.assertEqual(float(RendimentoHCConfig.get().meta_rendimento), 96.5)

    def test_novo_motivo_vira_campo_do_formulario(self):
        self.client.post(reverse("laboratorio:hc_cadastro_save"), {"tipo": "MOTIVO", "nome": "Bolha"})
        novo = HCCadastro.objects.get(tipo="MOTIVO", nome="Bolha")
        resp = self.client.get(reverse("laboratorio:hc_registro_create"))
        self.assertIn(f"rt_{novo.pk}", resp.context["form"].fields)

    def test_excluir_cadastro_em_uso_apenas_inativa(self):
        self.post_registro(teve_retrabalho=1, **{f"rt_{self.risco.pk}": 2})
        self.client.post(reverse("laboratorio:hc_cadastro_delete", args=[self.risco.pk]))
        self.risco.refresh_from_db()
        self.assertFalse(self.risco.ativo)
        livre = HCCadastro.objects.create(tipo="MOTIVO", nome="Temporario")
        self.client.post(reverse("laboratorio:hc_cadastro_delete", args=[livre.pk]))
        self.assertFalse(HCCadastro.objects.filter(pk=livre.pk).exists())

    def test_paginas_de_configuracao_abrem(self):
        for nome in ("hc_config_campos", "hc_config_cadastros", "hc_config_integracoes", "hc_registros", "hc_registro_create"):
            self.assertEqual(self.client.get(reverse(f"laboratorio:{nome}")).status_code, 200, nome)
        for aba in ("GARRA", "MATERIAL", "MOTIVO", "EQUIPE"):
            self.assertEqual(self.client.get(reverse("laboratorio:hc_config_cadastros"), {"aba": aba}).status_code, 200)


class DashboardTests(RendimentoHCBase):
    def setUp(self):
        super().setUp()
        # dia 5: 1º passo (lente nova) 100 lentes / 5 retrab; 2º passo 20 / 4
        self.post_registro(qtd_lentes=100, teve_retrabalho=1, **{f"rt_{self.risco.pk}": 3, f"rt_{self.verniz.pk}": 2,
                                                              f"pm_{self.orma.pk}": 100, f"rm_{self.orma.pk}": 5})
        self.post_registro(qtd_lentes=20, tipo_garra=self.retr1.pk, turno=self.turno2.pk, colaborador=self.bia.pk,
                           teve_retrabalho=1, **{f"rt_{self.risco.pk}": 4})

    def test_calculos_consolidado(self):
        dados = hc.consolidado(date(2026, 10, 1), date(2026, 10, 31))
        dia5 = next(l for l in dados["linhas"] if l["data"] == date(2026, 10, 5))
        self.assertEqual((dia5["produzidas"], dia5["retrabalho"]), (120, 9))
        self.assertEqual(dia5["rend1"], 95.0)
        self.assertEqual(dia5["rend2"], 80.0)
        self.assertEqual(dia5["rendimento"], 92.5)
        vazio = next(l for l in dados["linhas"] if l["data"] == date(2026, 10, 6))
        self.assertIsNone(vazio["rendimento"])

    def test_analises(self):
        diaria = hc.analise_diaria(date(2026, 10, 5))
        risco = next(m for m in diaria["motivos"] if m["nome"] == "Risco")
        self.assertEqual(risco["qtd"], 7)
        self.assertEqual({t["nome"] for t in diaria["turnos"]}, {"1º Turno", "2º Turno"})
        orma = next(m for m in diaria["materiais"] if m["nome"].startswith("Orma"))
        self.assertEqual((orma["prod"], orma["ret"], orma["rend"]), (100, 5, 95.0))
        motivo = hc.retrabalho_motivo(date(2026, 10, 1), date(2026, 10, 31))
        risco = next(l for l in motivo["linhas"] if l["nome"] == "Risco")
        self.assertEqual(risco["total"], 7)
        self.assertEqual([c["qtd"] for c in risco["turnos"]], [3, 4])
        colab = hc.analise_colaborador(f"id:{self.bia.pk}", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual((colab["total"]["prod"], colab["total"]["ret"], colab["total"]["rend"]), (20, 4, 80.0))
        ano = hc.retrabalho_ano(2026)
        self.assertEqual(ano["linhas"][9]["produzidas"], 120)
        mat = hc.rendimento_material(self.orma.pk, date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(mat["total"]["prod"], 100)

    def test_todas_as_abas_renderizam(self):
        url = reverse("laboratorio:hc_dashboard")
        extras = {"colaborador": {"colaborador": f"id:{self.ana.pk}"}, "material": {"material": self.orma.pk},
                  "diaria": {"data": "2026-10-05"}, "ano": {"ano": 2026}}
        for aba in ("consolidado", "diaria", "motivo", "colaborador", "ano", "material"):
            resp = self.client.get(url, {"aba": aba, **extras.get(aba, {})})
            self.assertEqual(resp.status_code, 200, aba)
        resp = self.client.get(url, {"aba": "consolidado"})
        self.assertContains(resp, "92,50%")  # rendimento geral do dia 5 (mês do último lançamento)

    def test_dashboard_vazio(self):
        RendimentoHCRegistro.objects.all().delete()
        self.assertEqual(self.client.get(reverse("laboratorio:hc_dashboard")).status_code, 200)
        self.assertEqual(self.client.get(reverse("laboratorio:hc_dashboard"), {"aba": "colaborador", "colaborador": "lixo"}).status_code, 200)


class IntegracaoTests(RendimentoHCBase):
    def test_sincronizacao_com_equipe_coating_e_rh(self):
        nova = _colab("10", "NOVA NO COATING")
        EquipeCoating.objects.create(colaborador=nova)
        inativa = _colab("11", "SAIU DA EMPRESA", ativo=False)
        EquipeHC.objects.create(colaborador=inativa)
        propostas = hc.propostas_sincronizacao()
        self.assertEqual(propostas["adicionar"], [nova])
        self.assertEqual([e.colaborador for e in propostas["inativos"]], [inativa])
        self.client.post(reverse("laboratorio:hc_sincronizar"))
        self.assertTrue(EquipeHC.objects.get(colaborador=nova).ativo)
        self.assertFalse(EquipeHC.objects.get(colaborador=inativa).ativo)
        self.assertEqual(hc.propostas_sincronizacao()["adicionar"], [])

    def test_vincular_nome_antigo_ao_colaborador(self):
        RendimentoHCRegistro.objects.create(data=date(2026, 10, 1), colaborador_nome="Aninha", qtd_lentes=10)
        self.assertEqual(hc.propostas_sincronizacao()["sem_vinculo"][0]["colaborador_nome"], "Aninha")
        self.client.post(reverse("laboratorio:hc_vincular_nome"), {"nome": "Aninha", "colaborador": self.ana.pk})
        self.assertEqual(RendimentoHCRegistro.objects.get().colaborador, self.ana)
        self.assertIn("Aninha", EquipeHC.objects.get(colaborador=self.ana).apelidos)

    def test_resolver_nome_por_prefixo_e_apelido(self):
        cands = hc.candidatos_colaborador()
        self.assertEqual(hc.resolver_colaborador("Ana Paula", cands), self.ana)
        self.assertEqual(hc.resolver_colaborador("ana paula souza", cands), self.ana)
        self.assertIsNone(hc.resolver_colaborador("Desconhecida", cands))
        membro = EquipeHC.objects.get(colaborador=self.bia)
        membro.apelidos = "Bibi"
        membro.save()
        self.assertEqual(hc.resolver_colaborador("bibi", hc.candidatos_colaborador()), self.bia)


class ImportacaoPlanilhaTests(RendimentoHCBase):
    def _planilha(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Respostas"
        ws.append(["Id", "Start time", "Completion time", "Email", "Name", "Turno", "Data", "Colaborador", "Número da\xa0Garra",
                   "Tipo de Garra", "Quantidade de Lentes", "Quantidade de Orma", "Quantidade de Poly", "Quantidade de Stylis",
                   "Quantidade de Poly BLUE UV", "Quantidade de Outros Fabricantes", "Teve Retrabalho ?",
                   "Quantidade Retrabalho de Orma", "Quantidade Retrabalho de Poly", "Quantidade Retrabalho de Stylis",
                   "Quantidade Retrabalho Poly BLUE UV", "Quantidade Retrabalho de Outros Fabricantes",
                   "Retrabalho\xa0Ponto de Verniz", "Retrabalho\xa0Escorrido de Verniz", "Retrabalho\xa0Risco"])
        ws.append([100, None, None, "anônima", None, "1º turno", date(2026, 10, 5), "Ana Paula", "3", "Lente Nova", 14, 10, 4,
                   None, None, None, "Sim", 2, None, None, None, None, 1, None, 1])
        ws.append([101, None, None, "anônima", None, "2º Turno", date(2026, 10, 5), "Desconhecido", "4", "1º Retrabalho", 10, 10,
                   None, None, None, None, "Não"])
        ws.append([102, None, None, "anônima", None, "2º Turno", date(2001, 3, 1), "Ana Paula", "4", "Lente Nova", 10])  # data inválida
        fd, caminho = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(caminho)
        self.addCleanup(os.remove, caminho)
        return caminho

    def test_importa_e_e_idempotente(self):
        caminho = self._planilha()
        call_command("importar_rendimento_hc", caminho, verbosity=0)
        self.assertEqual(RendimentoHCRegistro.objects.count(), 2)
        r = RendimentoHCRegistro.objects.get(chave_importacao="id:100")
        self.assertEqual((r.qtd_lentes, r.qtd_retrabalho, r.colaborador, r.turno), (14, 2, self.ana, self.turno1))
        self.assertEqual(r.tipo_garra, self.nova)
        self.assertEqual(r.itens.filter(tipo="PM").count(), 2)
        self.assertEqual(r.itens.get(tipo="RM").quantidade, 2)
        outro = RendimentoHCRegistro.objects.get(chave_importacao="id:101")
        self.assertIsNone(outro.colaborador)
        self.assertEqual(outro.colaborador_nome, "Desconhecido")
        call_command("importar_rendimento_hc", caminho, verbosity=0)
        self.assertEqual(RendimentoHCRegistro.objects.count(), 2)

    def test_dry_run_nao_grava(self):
        call_command("importar_rendimento_hc", self._planilha(), "--dry-run", verbosity=0)
        self.assertEqual(RendimentoHCRegistro.objects.count(), 0)


class ImportacaoPorTelaTests(ImportacaoPlanilhaTests):
    def _upload(self, url, caminho, **extra):
        with open(caminho, "rb") as f:
            return self.client.post(reverse(url), {"arquivo": f, **extra})

    def test_preview_sugere_o_mapeamento(self):
        resp = self._upload("laboratorio:hc_importar_preview", self._planilha())
        dados = resp.json()
        self.assertEqual(dados["aba"], "Respostas")
        self.assertEqual(dados["total"], 3)
        sug = dados["sugestao"]
        cab = dados["cabecalhos"]
        self.assertEqual(cab[sug["data"]], "Data")
        self.assertEqual(cab[sug["qtd_lentes"]], "Quantidade de Lentes")
        self.assertEqual(cab[sug[f"pm_{self.orma.pk}"]], "Quantidade de Orma")
        self.assertEqual(cab[sug[f"rt_{self.risco.pk}"]], "Retrabalho Risco")
        self.assertEqual(self.client.post(reverse("laboratorio:hc_importar_preview"), {}).status_code, 400)

    def test_importar_com_mapeamento_escolhido_pelo_usuario(self):
        caminho = self._planilha()
        sug = self._upload("laboratorio:hc_importar_preview", caminho).json()["sugestao"]
        mapa = {f"map_{k}": v for k, v in sug.items()}
        resp = self._upload("laboratorio:hc_importar", caminho, acao="simular", **mapa)
        self.assertEqual(resp.context["resultado"]["validas"], 2)
        self.assertEqual(RendimentoHCRegistro.objects.count(), 0)
        resp = self._upload("laboratorio:hc_importar", caminho, acao="importar", **mapa)
        self.assertEqual(resp.context["resultado"]["importadas"], 2)
        self.assertEqual(RendimentoHCRegistro.objects.get(chave_importacao="id:100").colaborador, self.ana)
        # coluna de lentes não relacionada -> erro e nada gravado
        sem_lentes = {k: v for k, v in mapa.items() if k != "map_qtd_lentes"}
        resp = self._upload("laboratorio:hc_importar", caminho, acao="importar", **sem_lentes)
        self.assertIsNone(resp.context["resultado"])
        self.assertEqual(RendimentoHCRegistro.objects.count(), 2)

    def test_pagina_abre(self):
        self.assertEqual(self.client.get(reverse("laboratorio:hc_importar")).status_code, 200)


class PermissoesHCTests(RendimentoHCBase):
    def test_todas_as_views_hc_estao_no_mapa_de_permissoes(self):
        from shared.permissions import get_view_permission_map

        mapa = get_view_permission_map()
        for nome in ("hc_dashboard", "hc_registros", "hc_registro_create", "hc_registro_update", "hc_registro_delete",
                     "hc_config_campos", "hc_campo_save", "hc_campo_mover", "hc_campo_delete", "hc_config_meta",
                     "hc_config_cadastros", "hc_cadastro_save", "hc_cadastro_delete", "hc_equipe_save", "hc_equipe_delete",
                     "hc_config_integracoes", "hc_sincronizar", "hc_vincular_nome", "hc_importar", "hc_importar_preview", "hc_modelo", "hc_exportar"):
            self.assertIn(f"laboratorio:{nome}", mapa, nome)

    def test_menu_mostra_itens_do_hc(self):
        resp = self.client.get(reverse("laboratorio:hc_dashboard"))
        self.assertContains(resp, reverse("laboratorio:hc_config_campos"))
        self.assertContains(resp, reverse("laboratorio:hc_registro_create"))


class DiaHoraHCTests(RendimentoHCBase):
    def _reg(self, hora, qtd, ret, garra, **kw):
        from datetime import datetime

        from django.utils import timezone

        r = RendimentoHCRegistro.objects.create(data=date(2026, 10, 5), qtd_lentes=qtd, qtd_retrabalho=ret, teve_retrabalho=bool(ret), tipo_garra=garra, **kw)
        if hora is not None:
            RendimentoHCRegistro.objects.filter(pk=r.pk).update(criado_em=timezone.make_aware(datetime(2026, 10, 5, hora, 15)))

    def test_agrupa_por_hora_com_passos(self):
        self._reg(9, 100, 5, self.nova)
        self._reg(9, 20, 4, self.retr1)
        self._reg(11, 50, 0, self.nova)
        self._reg(None, 30, 3, self.nova, origem=RendimentoHCRegistro.ORIGEM_IMPORT)
        d = hc.por_hora(date(2026, 10, 5))
        self.assertEqual([l["rotulo"] for l in d["linhas"]], ["09h", "10h", "11h"])
        h9, h10, h11 = d["linhas"]
        self.assertEqual((h9["produzidas"], h9["retrabalho"], h9["rend1"], h9["rend2"], h9["rendimento"]), (120, 9, 95.0, 80.0, 92.5))
        self.assertEqual((h10["produzidas"], h10["rendimento"]), (0, None))
        self.assertEqual(h11["rend1"], 100.0)
        self.assertEqual((d["sem_hora"]["produzidas"], d["total"]["produzidas"], d["total"]["retrabalho"]), (30, 200, 12))

    def test_tela_renderiza(self):
        url = reverse("laboratorio:hc_dia_hora")
        self.assertEqual(self.client.get(url).status_code, 200)
        self._reg(9, 100, 5, self.nova)
        resp = self.client.get(url, {"data": "2026-10-05"})
        self.assertContains(resp, "09h")
        self.assertContains(resp, "95,00%")
        self.assertEqual(self.client.get(url, {"data": "lixo"}).status_code, 200)

    def test_view_no_mapa_de_permissoes(self):
        from shared.permissions import get_view_permission_map

        self.assertIn("laboratorio:hc_dia_hora", get_view_permission_map())

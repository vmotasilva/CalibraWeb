"""Modelo de planilha (para importar) e exportação dos registros de Rendimento HC e Retrabalho CF.

Os cabeçalhos seguem o formato que o importador reconhece sozinho, então uma planilha exportada pode ser reimportada
(em outro ambiente, por exemplo) e o modelo vazio já vem pronto para ser preenchido.
"""
import io
from collections import defaultdict

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .hc_importacao import ALIAS_MATERIAL
from .models import (
    CFCadastro,
    CFCampoFormulario,
    HCCadastro,
    HCCampoFormulario,
    RendimentoHCItem,
    RetrabalhoCFItem,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ORIGEM_NOME = dict(CFCadastro.ORIGEM_CHOICES)


def _nome_curto_material(m):
    """Mesmo critério do importador para reconhecer a coluna do material."""
    for nome, aliases in ALIAS_MATERIAL.items():
        if m.nome == nome:
            return {"ORMA": "Orma", "POLY": "Poly", "STYLIS": "Stylis", "POLYBLUEUV": "Poly Blue UV",
                    "OUTROSFABRICANTES": "Outros Fabricantes"}[aliases[0]]
    return m.nome.split("/")[0].strip()


# --------------------------------------------------------------------------- colunas
def colunas_hc():
    """[(cabeçalho, função(registro, quantidades) -> valor)] na ordem da planilha original."""
    materiais = list(HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MATERIAL, ativo=True))
    motivos = list(HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MOTIVO, ativo=True))
    cols = [
        ("Id", lambda r, q: int(r.chave_importacao[3:]) if (r.chave_importacao or "").startswith("id:") else None),
        ("Data", lambda r, q: r.data),
        ("Hora", lambda r, q: r.hora.strftime("%H:%M") if r.hora else ""),
        ("Turno", lambda r, q: r.turno_nome),
        ("Colaborador", lambda r, q: r.colaborador_display),
        ("Número da Garra", lambda r, q: r.numero_garra),
        ("Tipo de Garra", lambda r, q: r.tipo_garra.nome if r.tipo_garra_id else ""),
        ("Tratamento", lambda r, q: r.tratamento.nome if r.tratamento_id else ""),
        ("Quantidade de Lentes", lambda r, q: r.qtd_lentes),
    ]
    cols += [(f"Quantidade de {_nome_curto_material(m)}", (lambda m: lambda r, q: q.get((RendimentoHCItem.PROD_MATERIAL, m.pk)))(m)) for m in materiais]
    cols.append(("Teve Retrabalho ?", lambda r, q: "Sim" if r.teve_retrabalho else "Não"))
    cols += [(f"Quantidade Retrabalho de {_nome_curto_material(m)}", (lambda m: lambda r, q: q.get((RendimentoHCItem.RET_MATERIAL, m.pk)))(m)) for m in materiais]
    cols += [(f"Retrabalho {m.nome}", (lambda m: lambda r, q: q.get((RendimentoHCItem.RET_MOTIVO, m.pk)))(m)) for m in motivos]
    cols += [(c.rotulo, (lambda c: lambda r, q: (r.extras or {}).get(c.chave))(c)) for c in HCCampoFormulario.objects.filter(sistema=False)]
    cols.append(("Observações", lambda r, q: r.observacoes))
    cols.append(("Retrabalho Total", lambda r, q: r.qtd_retrabalho))
    return cols


def colunas_cf():
    motivos = list(CFCadastro.objects.filter(tipo="MOTIVO", ativo=True))
    cols = [
        ("Id", lambda r, q: int(r.chave_importacao[3:]) if (r.chave_importacao or "").startswith("id:") else None),
        ("Data", lambda r, q: r.data),
        ("Hora", lambda r, q: r.hora.strftime("%H:%M") if r.hora else ""),
        ("Turno", lambda r, q: r.turno_nome),
        ("Número do Lote", lambda r, q: r.numero_lote),
        ("Total de Lentes no Lote", lambda r, q: r.total_lentes),
        ("Máquina", lambda r, q: r.maquina_nome),
        ("Tratamento", lambda r, q: r.tratamento_nome),
        ("Colaborador CF - Responsável Inspeção do Lote", lambda r, q: r.colaborador_cf_display),
        ("Necessidade Retrabalho?", lambda r, q: "Sim" if r.teve_retrabalho else "Não"),
    ]
    cols += [(f"Retrabalho {m.grupo} - {m.grupo} - {m.nome}", (lambda m: lambda r, q: q.get(m.pk))(m)) for m in motivos]
    cols += [
        ("Lote Reprovado?", lambda r, q: "Sim" if r.reprovado else "Não"),
        ("Motivo da Reprovação", lambda r, q: r.motivo_reprovacao),
        ("Colaborador AR - Responsável Montagem Lote", lambda r, q: r.colaborador_ar_display),
    ]
    cols += [(c.rotulo, (lambda c: lambda r, q: (r.extras or {}).get(c.chave))(c)) for c in CFCampoFormulario.objects.filter(sistema=False)]
    cols += [("Observações", lambda r, q: r.observacoes), ("Retrabalho Total", lambda r, q: r.qtd_retrabalho)]
    return cols


# --------------------------------------------------------------------------- instruções do modelo
INSTRUCOES = {
    "hc": [
        ("Como usar", "Preencha a aba 'Respostas' (um lançamento por linha, uma linha por garra) e importe em HMC > HC > Importar Planilha."),
        ("Data", "Obrigatória. Formato dd/mm/aaaa."),
        ("Quantidade de Lentes", "Obrigatória. Total de lentes da garra."),
        ("Turno / Tipo de Garra / Tratamento", "Use exatamente os nomes cadastrados (Regras de Turno, Cadastros base e Tratamentos)."),
        ("Colaborador", "Nome do colaborador (o sistema vincula ao RH; apelidos da Equipe HC também são reconhecidos)."),
        ("Quantidade de <material> / Quantidade Retrabalho de <material>", "Lentes produzidas e retrabalhadas por material. Deixe em branco o que não se aplica."),
        ("Retrabalho <motivo>", "Quantidade retrabalhada por motivo. O total de retrabalho é a soma desses campos."),
        ("Id", "Opcional. Se informado, evita duplicar o lançamento ao reimportar a mesma planilha."),
        ("Colunas extras", "Colunas 'Teve Retrabalho ?' e 'Retrabalho Total' são informativas (calculadas pelo sistema)."),
    ],
    "cf": [
        ("Como usar", "Preencha a aba 'Respostas' (uma linha por lote inspecionado) e importe em HMC > Controle Final > Importar Planilha."),
        ("Data", "Obrigatória. Formato dd/mm/aaaa."),
        ("Total de Lentes no Lote", "Obrigatória."),
        ("Turno / Máquina / Tratamento", "Use os nomes cadastrados; nomes diferentes ficam como texto até serem vinculados em Integrações."),
        ("Colaborador CF / Colaborador AR", "Inspetor do lote e responsável pela montagem."),
        ("Retrabalho <origem> - <origem> - <motivo>", "Quantidade retrabalhada por motivo (AR, HC ou CF). O total de retrabalho é a soma desses campos."),
        ("Lote Reprovado? / Motivo da Reprovação", "Sim/Não e, se reprovado, o motivo cadastrado."),
        ("Id", "Opcional. Se informado, evita duplicar o lote ao reimportar a mesma planilha."),
        ("Colunas extras", "'Necessidade Retrabalho?' e 'Retrabalho Total' são informativas (calculadas pelo sistema)."),
    ],
}


# --------------------------------------------------------------------------- geração
def _resposta(wb, nome_arquivo):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'
    return resp


def _valor(v):
    if v is None or isinstance(v, (int, float)):
        return v
    return v if hasattr(v, "year") else str(v)


def _escrever(ws, wb, cabecalhos, linhas):
    estilo_fonte, estilo_fundo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14284B")
    for i in range(len(cabecalhos)):
        ws.column_dimensions[_letra(i + 1)].width = 16
    cab = []
    for h in cabecalhos:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.fill = estilo_fonte, estilo_fundo
        cab.append(c)
    ws.append(cab)
    ws.freeze_panes = "A2"
    for linha in linhas:
        ws.append(linha)


def _letra(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _planilha(cols, linhas, instrucoes):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Respostas")
    _escrever(ws, wb, [h for h, _f in cols], linhas)
    ws2 = wb.create_sheet("Instruções")
    ws2.column_dimensions["A"].width = 48
    ws2.column_dimensions["B"].width = 110
    for par in [("Campo", "Orientação")] + instrucoes:
        ws2.append(list(par))
    return wb


def modelo(kind):
    cols = colunas_hc() if kind == "hc" else colunas_cf()
    nome = "modelo_rendimento_hc.xlsx" if kind == "hc" else "modelo_retrabalho_cf.xlsx"
    return _resposta(_planilha(cols, [], INSTRUCOES[kind]), nome)


def exportar(kind, qs):
    """Exporta os registros (queryset já filtrado) no formato do importador."""
    if kind == "hc":
        cols = colunas_hc()
        quantidades = defaultdict(dict)
        for i in RendimentoHCItem.objects.filter(registro__in=qs).values_list("registro_id", "tipo", "cadastro_id", "quantidade"):
            quantidades[i[0]][(i[1], i[2])] = i[3]
        nome = "rendimento_hc"
    else:
        cols = colunas_cf()
        quantidades = defaultdict(dict)
        for i in RetrabalhoCFItem.objects.filter(registro__in=qs).values_list("registro_id", "motivo_id", "quantidade"):
            quantidades[i[0]][i[1]] = i[2]
        nome = "retrabalho_cf"
    linhas = ([_valor(f(r, quantidades.get(r.pk, {}))) for _h, f in cols] for r in qs.order_by("data", "id").iterator(chunk_size=1000))
    carimbo = timezone.localtime().strftime("%Y%m%d_%H%M")
    return _resposta(_planilha(cols, linhas, INSTRUCOES[kind]), f"{nome}_{carimbo}.xlsx")

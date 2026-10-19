"""Importação de planilhas (xlsx/csv) para os lançamentos de Rendimento HC, com mapeamento coluna -> campo do sistema."""
import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime, time

import openpyxl
from django.db import transaction

from . import hc_rendimento as hc
from .models import (
    HCCadastro,
    HCCampoFormulario,
    RegraTurnoCoating,
    RendimentoHCItem,
    RendimentoHCRegistro,
    TratamentoAntiReflexo,
)

ANO_MIN = 2020
LOTE = 500

# (chave, rótulo, obrigatório, palavras que identificam o cabeçalho — comparação por igualdade, sem espaços/acentos)
CAMPOS_FIXOS = [
    ("id", "Id do registro (evita duplicar ao reimportar)", False, ["ID"]),
    ("data", "Data", True, ["DATA"]),
    ("data_ref", "Hora / data e hora de preenchimento", False, ["HORA", "STARTTIME", "COMPLETIONTIME", "HORADOLANCAMENTO", "CARIMBODEDATAHORA"]),
    ("turno", "Turno", False, ["TURNO"]),
    ("colaborador", "Colaborador", False, ["COLABORADOR", "NOME"]),
    ("numero_garra", "Número da garra", False, ["NUMERODAGARRA", "GARRA"]),
    ("tipo_garra", "Tipo de garra", False, ["TIPODEGARRA"]),
    ("tratamento", "Tratamento", False, ["TRATAMENTO"]),
    ("qtd_lentes", "Quantidade de lentes", True, ["QUANTIDADEDELENTES", "QUANTIDADELENTES"]),
    ("observacoes", "Observações", False, ["OBSERVACOES", "OBSERVACAO"]),
]
ALIAS_MATERIAL = {
    "Orma / Orgânico / CR39": ["ORMA"],
    "Airwear / TFL / Poly": ["POLY"],
    "Stylis 1.67 / Stylis 1.74": ["STYLIS"],
    "Poly Blue UV": ["POLYBLUEUV"],
    "Lista Negra / Terceiros (Outros Fabricantes)": ["OUTROSFABRICANTES", "OUTROS"],
}


def _compacto(texto):
    return re.sub(r"[^A-Z0-9]", "", hc.norm(str(texto or "")))


def alvos():
    """Todos os campos do sistema que podem receber uma coluna: [{chave, rotulo, obrigatorio, grupo}]."""
    out = [{"chave": k, "rotulo": r, "obrigatorio": o, "grupo": "Dados do lançamento"} for k, r, o, _a in CAMPOS_FIXOS]
    for m in HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MATERIAL, ativo=True):
        out.append({"chave": f"pm_{m.pk}", "rotulo": f"Lentes produzidas — {m.nome}", "obrigatorio": False, "grupo": "Produção por material"})
    for m in HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MATERIAL, ativo=True):
        out.append({"chave": f"rm_{m.pk}", "rotulo": f"Retrabalho — {m.nome}", "obrigatorio": False, "grupo": "Retrabalho por material"})
    for m in HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MOTIVO, ativo=True):
        out.append({"chave": f"rt_{m.pk}", "rotulo": f"Retrabalho — {m.nome}", "obrigatorio": False, "grupo": "Retrabalho por motivo"})
    for c in HCCampoFormulario.objects.filter(sistema=False):
        out.append({"chave": f"x_{c.chave}", "rotulo": c.rotulo, "obrigatorio": False, "grupo": "Perguntas personalizadas"})
    return out


# --------------------------------------------------------------------------- leitura
def ler_planilha(arquivo, aba=None):
    """Lê xlsx/csv. Devolve {abas, aba, cabecalhos, linhas(list de list)}."""
    nome = (getattr(arquivo, "name", "") or "").lower()
    conteudo = arquivo.read()
    if nome.endswith(".csv"):
        texto = conteudo.decode("utf-8-sig", errors="replace")
        dialeto = csv.excel_tab if "\t" in texto[:2000] else (csv.excel if texto[:2000].count(",") > texto[:2000].count(";") else _PontoVirgula)
        linhas = [list(r) for r in csv.reader(io.StringIO(texto), dialeto)]
        abas, aba = ["CSV"], "CSV"
    else:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        try:
            abas = wb.sheetnames
            aba = aba if aba in abas else ("Respostas" if "Respostas" in abas else abas[0])
            linhas = [list(r) for r in wb[aba].iter_rows(values_only=True)]
        finally:
            wb.close()
    linhas = [r for r in linhas if r and any(v not in (None, "") for v in r)]
    if not linhas:
        raise ValueError("A planilha está vazia.")
    cab = [str(c).strip() if c not in (None, "") else f"(coluna {i + 1})" for i, c in enumerate(linhas[0])]
    return {"abas": abas, "aba": aba, "cabecalhos": cab, "linhas": linhas[1:]}


class _PontoVirgula(csv.excel):
    delimiter = ";"


def sugerir_mapa(cabecalhos):
    """{chave do campo: índice da coluna} sugerido por nome de cabeçalho."""
    comp = [_compacto(c) for c in cabecalhos]
    usados, mapa = set(), {}

    def achar(opcoes):
        for op in opcoes:
            for i, c in enumerate(comp):
                if c == op and i not in usados:
                    return i
        return None

    def marcar(chave, opcoes):
        i = achar(opcoes)
        if i is not None:
            mapa[chave] = i
            usados.add(i)

    for chave, _r, _o, aliases in CAMPOS_FIXOS:
        marcar(chave, aliases)
    materiais = HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MATERIAL, ativo=True)
    for m in materiais:
        aliases = ALIAS_MATERIAL.get(m.nome) or [_compacto(m.nome.split("/")[0])]
        marcar(f"pm_{m.pk}", [f"QUANTIDADEDE{a}" for a in aliases] + [f"QUANTIDADE{a}" for a in aliases])
    for m in materiais:
        aliases = ALIAS_MATERIAL.get(m.nome) or [_compacto(m.nome.split("/")[0])]
        marcar(f"rm_{m.pk}", [f"QUANTIDADERETRABALHODE{a}" for a in aliases] + [f"QUANTIDADERETRABALHO{a}" for a in aliases])
    for m in HCCadastro.objects.filter(tipo=HCCadastro.TIPO_MOTIVO, ativo=True):
        marcar(f"rt_{m.pk}", ["RETRABALHO" + _compacto(m.nome)])
    for c in HCCampoFormulario.objects.filter(sistema=False):
        marcar(f"x_{c.chave}", [_compacto(c.rotulo)])
    return mapa


# --------------------------------------------------------------------------- conversões
def _int(v):
    try:
        return max(int(float(str(v).replace(",", "."))), 0)
    except (TypeError, ValueError):
        return 0


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _hora(v):
    """Hora (HH:MM) de um valor de planilha (datetime, time ou texto 'dd/mm/aaaa HH:MM[:SS]'); meia-noite = sem hora."""
    if isinstance(v, datetime):
        t = v.time()
    elif isinstance(v, time):
        t = v
    else:
        m = re.search(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)(?::[0-5]\d)?(?!\d)", str(v or ""))
        t = time(int(m.group(1)), int(m.group(2))) if m else None
    if t is None:
        return None
    t = t.replace(second=0, microsecond=0)
    return t if t != time() else None


def _atualizar_horas(modelo, horas):
    """Preenche `hora` de registros já importados (reimportação) que ainda não a tinham. Devolve quantos."""
    if not horas:
        return 0
    n = 0
    for ini in range(0, len(horas), LOTE):
        parte = dict(list(horas.items())[ini: ini + LOTE])
        regs = list(modelo.objects.filter(chave_importacao__in=parte, hora__isnull=True))
        for r in regs:
            r.hora = parte[r.chave_importacao]
        modelo.objects.bulk_update(regs, ["hora"])
        n += len(regs)
    return n


# --------------------------------------------------------------------------- importação
def importar(linhas, mapa, a_partir=None, simular=False, usuario=None):
    """Importa as linhas conforme `mapa` ({chave: índice}). Devolve um dicionário de estatísticas."""
    hc.garantir_padroes()
    mapa = {k: int(v) for k, v in mapa.items() if v not in (None, "")}
    if "data" not in mapa or "qtd_lentes" not in mapa:
        raise ValueError("Relacione ao menos as colunas de Data e de Quantidade de lentes.")

    def cel(row, chave):
        i = mapa.get(chave)
        return row[i] if i is not None and i < len(row) else None

    garras = {hc.norm(c.nome): c for c in HCCadastro.objects.filter(tipo=HCCadastro.TIPO_GARRA)}
    trats = {hc.norm(t.nome): t for t in TratamentoAntiReflexo.objects.all()}
    turnos = {hc.norm(t.nome): t for t in RegraTurnoCoating.objects.all()}
    turnos_num = {}
    for t in RegraTurnoCoating.objects.all():
        m = re.search(r"[1-9]", t.nome)
        if m:
            turnos_num.setdefault(int(m.group()), t)
    candidatos = hc.candidatos_colaborador()
    cache_colab, cache_nome = {}, {}
    pm = [(k, int(k[3:])) for k in mapa if k.startswith("pm_")]
    rm = [(k, int(k[3:])) for k in mapa if k.startswith("rm_")]
    rt = [(k, int(k[3:])) for k in mapa if k.startswith("rt_")]
    extras_keys = [k for k in mapa if k.startswith("x_")]

    existentes = set(RendimentoHCRegistro.objects.exclude(chave_importacao=None).values_list("chave_importacao", flat=True))
    hoje = date.today()
    novos, pendentes = [], []
    st = Counter()
    sem_garra, sem_turno, sem_trat = Counter(), Counter(), Counter()
    horas_existentes = {}

    for row in linhas:
        dia = _data(cel(row, "data"))
        if not dia or dia.year < ANO_MIN or dia > date(hoje.year + 1, 12, 31):
            st["ignoradas"] += 1
            continue
        if a_partir and dia < a_partir:
            st["antes_do_corte"] += 1
            continue
        id_ext = _int(cel(row, "id")) if "id" in mapa else 0
        chave = f"id:{id_ext}" if id_ext else "sha:" + hashlib.md5(repr(tuple(row)).encode()).hexdigest()[:24]
        hora = _hora(cel(row, "data_ref")) if "data_ref" in mapa else None
        hora = hora or _hora(cel(row, "data"))
        if chave in existentes:
            st["repetidas"] += 1
            if hora:
                horas_existentes[chave] = hora
            continue
        existentes.add(chave)

        turno_txt = str(cel(row, "turno") or "").strip()
        turno = turnos.get(hc.norm(turno_txt))
        if not turno and turno_txt:
            m = re.search(r"[1-9]", turno_txt)
            turno = turnos_num.get(int(m.group())) if m else None
        if turno_txt and not turno:
            sem_turno[turno_txt] += 1
        nome = str(cel(row, "colaborador") or "").strip()
        if nome not in cache_colab:
            cache_colab[nome] = hc.resolver_colaborador(nome, candidatos)
        colab = cache_colab[nome]
        garra_txt = str(cel(row, "tipo_garra") or "").strip()
        garra = garras.get(hc.norm(garra_txt))
        if garra_txt and not garra:
            sem_garra[garra_txt] += 1
        trat_txt = str(cel(row, "tratamento") or "").strip()
        trat = trats.get(hc.norm(trat_txt))
        if trat_txt and not trat:
            sem_trat[trat_txt] += 1

        motivos = {cid: _int(cel(row, k)) for k, cid in rt}
        total = sum(motivos.values())
        itens = [(RendimentoHCItem.PROD_MATERIAL, cid, _int(cel(row, k))) for k, cid in pm if _int(cel(row, k))]
        itens += [(RendimentoHCItem.RET_MATERIAL, cid, _int(cel(row, k))) for k, cid in rm if _int(cel(row, k))]
        itens += [(RendimentoHCItem.RET_MOTIVO, cid, q) for cid, q in motivos.items() if q]
        por_material = sum(q for t, _c, q in itens if t == RendimentoHCItem.RET_MATERIAL)
        total = total or por_material
        qtd = _int(cel(row, "qtd_lentes"))
        extras = {}
        for k in extras_keys:
            v = cel(row, k)
            if v not in (None, ""):
                extras[k[2:]] = v.isoformat() if isinstance(v, (date, datetime)) else v
        numero = cel(row, "numero_garra")
        novos.append(RendimentoHCRegistro(
            data=dia, hora=hora, turno=turno, turno_nome=turno.nome if turno else turno_txt[:50],
            colaborador=colab, colaborador_nome=colab.nome_completo if colab else nome[:150],
            numero_garra=("" if numero is None else str(int(numero)) if isinstance(numero, float) and numero.is_integer() else str(numero)).strip()[:10],
            tipo_garra=garra, tratamento=trat, qtd_lentes=qtd, teve_retrabalho=total > 0, qtd_retrabalho=total,
            observacoes=str(cel(row, "observacoes") or "")[:2000], extras=extras,
            origem=RendimentoHCRegistro.ORIGEM_IMPORT, chave_importacao=chave, usuario=usuario,
        ))
        pendentes.append(itens)
        if qtd == 0 or total > qtd:
            st["a_revisar"] += 1

    resultado = {
        "validas": len(novos), "repetidas": st["repetidas"], "ignoradas": st["ignoradas"], "antes_do_corte": st["antes_do_corte"],
        "a_revisar": st["a_revisar"], "com_colaborador": sum(1 for r in novos if r.colaborador_id),
        "sem_colaborador": sorted({r.colaborador_nome for r in novos if not r.colaborador_id and r.colaborador_nome})[:30],
        "turno_sem_regra": dict(sem_turno), "garra_sem_cadastro": dict(sem_garra), "tratamento_sem_cadastro": dict(sem_trat),
        "simulado": simular, "importadas": 0,
    }
    resultado["horas_preenchidas"] = len(horas_existentes) if simular else _atualizar_horas(RendimentoHCRegistro, horas_existentes)
    if simular or not novos:
        return resultado
    with transaction.atomic():
        for ini in range(0, len(novos), LOTE):
            lote = novos[ini: ini + LOTE]
            RendimentoHCRegistro.objects.bulk_create(lote)
            ids = {r.chave_importacao: r.pk for r in RendimentoHCRegistro.objects.filter(chave_importacao__in=[r.chave_importacao for r in lote])}
            itens = []
            for reg, pend in zip(lote, pendentes[ini: ini + LOTE]):
                pk = reg.pk or ids.get(reg.chave_importacao)
                itens += [RendimentoHCItem(registro_id=pk, tipo=t, cadastro_id=cid, quantidade=q) for t, cid, q in pend if pk]
            RendimentoHCItem.objects.bulk_create(itens)
    resultado["importadas"] = len(novos)
    return resultado

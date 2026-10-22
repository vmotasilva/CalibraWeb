"""Importação da estrutura dos formulários de PCP a partir das planilhas (.xlsx) em uso.

Lê só a estrutura (itens, especificações, métodos, frequências, planos de reação); as colunas de
data x turno com os lançamentos antigos são ignoradas. Cada aba vira uma Página e a revisão importada
nasce como rascunho (ou vigente, se pedido).
"""
import re
import unicodedata
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from organization.models import Setor
from pcp.models import Campo, Formulario, FormularioVersao, Frequencia, ListaOpcoes, OpcaoLista, Pagina, Turno

ABAS_IGNORADAS = {'dados', 'indicador pcp', 'indicador', 'folha1', 'planilha1', 'planilha2', 'indice'}
MAX_LINHAS = 700
MAX_COLUNAS = 14

RE_NOME_ARQUIVO = re.compile(r'(FOR\.\d+)\.(R\d+)\s*-\s*(.+?)\s*\.xlsx?$', re.I)
RE_FAIXA = re.compile(r'^\s*(-?\d+(?:[.,]\d+)?)\s*(?:-|–|a|até)\s*(-?\d+(?:[.,]\d+)?)\s*(.*)$', re.I)
RE_MAIS_MENOS = re.compile(r'^\s*(-?\d+(?:[.,]\d+)?)\s*([^\d±+]{0,10}?)\s*(?:±|\+\s*/\s*-)\s*(\d+(?:[.,]\d+)?)\s*(.{0,20})$')
RE_NUM = re.compile(r'^\s*-?\d+(?:[.,]\d+)?\s*$')


def _sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def _chave(s):
    """Chave de comparação: minúsculas, sem acento, só letras e dígitos."""
    return re.sub(r'[^a-z0-9]', '', _sem_acento(str(s)).lower())


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else ('%.6f' % v).rstrip('0').rstrip('.')
    if isinstance(v, (int, str)):
        return ' '.join(str(v).split())
    return ''


def _dec(s):
    try:
        return Decimal(str(s).strip().replace(',', '.'))
    except InvalidOperation:
        return None


def parse_arquivo(nome):
    """'FOR.080.R15 - Formulário Surfaçagem - PCP ... .xlsx' -> (codigo, revisao, titulo)."""
    m = RE_NOME_ARQUIVO.search(nome or '')
    if not m:
        return None
    return m.group(1).upper(), m.group(2).upper(), m.group(3).strip()


def parse_faixa(texto):
    """'10 - 27 ° C' -> (Decimal('10'), Decimal('27'), '° C'); None se não for uma faixa numérica simples."""
    m = RE_FAIXA.match(texto or '')
    if not m:
        return None
    a, b = _dec(m.group(1)), _dec(m.group(2))
    if a is None or b is None or a > b:
        return None
    resto = m.group(3).strip()
    return a, b, resto if len(resto) <= 20 else ''


def parse_mais_menos(texto):
    """'45°C ± 2°C' -> (Decimal('43'), Decimal('47'), '°C'); None se o formato não for 'valor ± tolerância'."""
    m = RE_MAIS_MENOS.match(texto or '')
    if not m:
        return None
    alvo, tol = _dec(m.group(1)), _dec(m.group(3))
    if alvo is None or tol is None:
        return None
    return alvo - tol, alvo + tol, (m.group(2) or m.group(4)).strip()


FREQUENCIAS_CANONICAS = [
    ('A cada turno', 'turno', {'turno', 'a cada turno', 'cada turno', 'pelo menos a cada turno', 'pelo menos uma vez por turno'}),
    ('Diária', 'diaria', {'diario', 'diariamente', 'dia', 'pelo menos diariamente'}),
    ('Semanal', 'semanal', {'semanal', 'semanalmente', 'semana', 'pelo menos semanalmente'}),
    ('Quinzenal', 'quinzenal', {'quinzenal', 'a cada 2 semanas', 'pelo menos uma vez a cada 2 semanas'}),
    ('Mensal', 'mensal', {'mensal', 'mensalmente', 'monthly', 'pelo menos mensalmente', 'uma vez por mes'}),
    ('Trimestral', 'trimestral', {'trimestral'}),
    ('Semestral', 'semestral', {'semestral', 'a cada 6 meses'}),
    ('Anual', 'anual', {'anual', 'anualmente'}),
]
FREQUENCIA_EVENTO = ('Sob demanda / por evento', 'evento')


def frequencia_canonica(texto):
    """Texto livre da planilha -> (nome, periodicidade) do cadastro. Textos específicos viram 'Sob demanda / por evento'."""
    t = re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', _sem_acento(texto).lower())).strip()
    for nome, per, sinonimos in FREQUENCIAS_CANONICAS:
        if t in sinonimos:
            return nome, per
    return FREQUENCIA_EVENTO


class _Cadastros:
    """Cache de turnos, listas e frequências usados durante a importação."""

    def __init__(self):
        self.freq = {_chave(f.nome): f for f in Frequencia.objects.all()}
        self.criadas_freq = 0
        self.listas = {}
        self._garantir_turnos()
        self.listas['ok'] = self._lista('Ok / Nok', [('Ok', True), ('Nok', False), ('Manut.', True), ('Máq. parada', True)])
        self.listas['real'] = self._lista('Realizado / Não realizado', [
            ('Realizado', True), ('Não realizado', False), ('Manut.', True), ('Máq. parada', True)])

    @staticmethod
    def _garantir_turnos():
        if not Turno.objects.exists():
            for i, sigla in enumerate(('T1', 'T2', 'T3'), 1):
                Turno.objects.create(nome=f'Turno {i}', sigla=sigla, ordem=i)

    @staticmethod
    def _lista(nome, opcoes):
        lista, _ = ListaOpcoes.objects.get_or_create(nome=nome)
        existentes = set(lista.opcoes.values_list('valor', flat=True))
        for ordem, (valor, conforme) in enumerate(opcoes):
            if valor not in existentes:
                OpcaoLista.objects.create(lista=lista, valor=valor, conforme=conforme, ordem=ordem)
        return lista

    def frequencia(self, texto):
        """(Frequencia do cadastro, detalhe) para o texto da planilha; o detalhe guarda o texto original se diferente."""
        texto = _txt(texto)
        if not texto or texto == '0' or texto.startswith('#'):
            return None, ''
        nome, per = frequencia_canonica(texto)
        k = _chave(nome)
        if k not in self.freq:
            self.freq[k] = Frequencia.objects.create(nome=nome, periodicidade=per, ordem=len(self.freq) + 1)
            self.criadas_freq += 1
        detalhe = '' if _chave(texto) == k else texto[:400]
        return self.freq[k], detalhe


# --------------------------------------------------------------------------------------
# Leitura das abas
# --------------------------------------------------------------------------------------

def _achar_cabecalho(linhas):
    """Retorna (indice_da_linha, layout, {campo: coluna}) ou None se a aba não tem um cabeçalho reconhecido."""
    for i, linha in enumerate(linhas[:12]):
        h = [_chave(c) for c in linha]

        def col(*chaves, exato=False, inicio=False):
            for ch in chaves:
                for j, v in enumerate(h):
                    if (v == ch) if exato else (v.startswith(ch) if inicio else ch in v):
                        return j
            return None

        if 'receita' in h and col('razao') is not None:
            return i, 'receita', dict(resp=col('resp'), freq=col('frequ'), etapa=col('receita', exato=True),
                                      descricao=col('razao'), spec=col('especifica'), tol=col('toleran'))
        if col('passo', exato=True) is not None and col('target', exato=True) is not None:
            return i, 'maquina', dict(resp=col('resp'), etapa=col('passo', exato=True), dispositivo=col('dispositivo', exato=True),
                                      descricao=col('comoverificar', inicio=True), target=col('target', exato=True),
                                      unidade=col('unidade', exato=True), minimo=col('toleranciaminima'),
                                      maximo=col('toleranciamaxima'), metodo=col('metododeverificacao'), freq=col('frequ'))
        if col('etapadoprocesso') is not None or (col('descricaodaoperacao') is not None and col('produto', inicio=True) is not None):
            p = col('processo', inicio=True)
            return i, 'etapa', dict(resp=col('responsavel'), etapa=col('etapadoprocesso'), operacao=col('descricaodaoperacao'),
                                    maquina=col('maquinaouequipamento'), n=col('n', exato=True) if col('n', exato=True) is not None else col('no', exato=True),
                                    produto=col('produto', inicio=True), processo=p, spec=(p + 1) if p is not None else None,
                                    metodo=col('dispositivoou'), freq=col('frequ'))
        if col('processoetapa') is not None:
            return i, 'processos', dict(resp=col('resp'), freq=col('frequ'), etapa=col('processoetapa'), dispositivo=col('maquinadispositivo'),
                                        descricao=col('parametro', exato=True), spec=col('especifica'), tol=col('toleran'),
                                        metodo=col('verificar', exato=True))
        if col('item', exato=True) is not None:
            return i, 'item', dict(item=col('item', exato=True), resp=col('resp'), descricao=col('oquever', 'verificar'),
                                   spec=col('especifica'), metodo=col('metodo', 'como'), freq=col('frequ'), plano=col('plano'))
    return None


def _linhas_campos(linhas, i_cab, layout, m):
    """Converte as linhas de dados em dicts de Campo. Estado (etapa/receita) é propagado para baixo."""
    def v(linha, nome):
        j = m.get(nome)
        return _txt(linha[j]) if j is not None and j < len(linha) else ''

    ultimo = {'etapa': ''}
    saida = []
    for linha in linhas[i_cab + 1:]:
        if any(_chave(c) in ('t1', 't2', 't3') for c in linha):
            continue  # segunda linha do cabeçalho (T1/T2/T3)
        c = dict(item='', etapa='', descricao='', especificacao='', metodo='', plano_reacao='', responsavel=v(linha, 'resp'),
                 frequencia=v(linha, 'freq'), tipo=None, minimo=None, maximo=None, unidade='')
        if layout == 'item':
            c.update(item=v(linha, 'item'), descricao=v(linha, 'descricao'), especificacao=v(linha, 'spec'),
                     metodo=v(linha, 'metodo'), plano_reacao=v(linha, 'plano'))
        elif layout == 'etapa':
            etapa_n, operacao, maquina = v(linha, 'etapa'), v(linha, 'operacao'), v(linha, 'maquina')
            if etapa_n or operacao:
                ultimo['etapa'] = ' · '.join(p for p in (f'{etapa_n}. {operacao}'.strip('. ') if etapa_n else operacao, maquina) if p)
            produto, processo = v(linha, 'produto'), v(linha, 'processo')
            if not (produto or processo):
                continue
            c.update(item=f'{etapa_n}.{v(linha, "n")}'.strip('.') if etapa_n else v(linha, 'n'), etapa=ultimo['etapa'],
                     descricao=f'{produto}: {processo}' if produto and processo else (produto or processo),
                     especificacao=v(linha, 'spec'), metodo=v(linha, 'metodo'))
        elif layout == 'receita':
            if v(linha, 'etapa'):
                ultimo['etapa'] = v(linha, 'etapa')
            razao = v(linha, 'descricao')
            if not razao:
                continue
            c.update(etapa=ultimo['etapa'], descricao=razao, especificacao=v(linha, 'spec'), tipo=Campo.NUMERO)
            faixa = re.match(r'^\s*(-?\d+(?:[.,]\d+)?)\s*-\s*(-?\d+(?:[.,]\d+)?)\s*$', v(linha, 'tol'))
            if faixa:
                c.update(minimo=_dec(faixa.group(1)), maximo=_dec(faixa.group(2)))
            else:
                c['especificacao'] = f"{c['especificacao']} (tolerância: {v(linha, 'tol')})" if v(linha, 'tol') else c['especificacao']
        elif layout == 'processos':
            if v(linha, 'etapa'):
                ultimo['etapa'] = v(linha, 'etapa')
            parametro = v(linha, 'descricao')
            if not parametro:
                continue
            disp, spec, tol = v(linha, 'dispositivo'), v(linha, 'spec'), v(linha, 'tol')
            c.update(etapa=ultimo['etapa'], descricao=f'{disp}: {parametro}' if disp else parametro,
                     especificacao=f'{spec} (tolerância: {tol})' if spec and tol else spec, metodo=v(linha, 'metodo'),
                     tipo=Campo.NUMERO if RE_NUM.match(spec) else Campo.TEXTO)
        elif layout == 'maquina':
            disp, como = v(linha, 'dispositivo'), v(linha, 'descricao')
            if not (disp or como):
                continue
            target, unidade = v(linha, 'target'), v(linha, 'unidade')
            c.update(etapa=v(linha, 'etapa'), descricao=f'{disp}: {como}' if disp and como else (disp or como),
                     especificacao=f'Target {target} {unidade}'.strip() if target else '', metodo=v(linha, 'metodo'),
                     unidade=unidade[:20], tipo=Campo.NUMERO if (_dec(v(linha, 'minimo')) is not None or _dec(v(linha, 'maximo')) is not None) else None,
                     minimo=_dec(v(linha, 'minimo')), maximo=_dec(v(linha, 'maximo')))
        if not c['descricao']:
            continue
        saida.append(c)
    return saida


def _tipo_e_limites(c, cad):
    """Define tipo/lista/limites quando a planilha não diz: faixa numérica na especificação vira número."""
    if c['tipo'] == Campo.NUMERO or c['tipo'] == Campo.TEXTO:
        return c['tipo'], None
    faixa = parse_faixa(c['especificacao'])
    if faixa:
        c['minimo'], c['maximo'] = faixa[0], faixa[1]
        c['unidade'] = c['unidade'] or faixa[2]
        return Campo.NUMERO, None
    pm = parse_mais_menos(c['especificacao'])
    if pm:
        c['minimo'], c['maximo'] = pm[0], pm[1]
        c['unidade'] = c['unidade'] or pm[2]
        return Campo.NUMERO, None
    limpeza = _chave(c['descricao']).startswith(('limpeza', 'limpar'))
    return Campo.LISTA, cad.listas['real' if limpeza else 'ok']


def _grupos_do_indice(wb):
    """{chave_da_aba: grupo} lido da aba Índice (cabeçalhos de coluna + nomes das páginas abaixo)."""
    nome = next((n for n in wb.sheetnames if _chave(n) == 'indice'), None)
    if not nome:
        return {}
    linhas = list(wb[nome].iter_rows(min_row=1, max_row=60, max_col=24, values_only=True))
    cab = None
    for i, linha in enumerate(linhas):
        cel = {j: _txt(c) for j, c in enumerate(linha) if _txt(c) and _txt(c) != '\xa0'}
        if len(cel) >= 4 and not any(re.match(r'(c[oó]digo|revis|data|p[aá]gina|[ií]ndice geral)', t, re.I) for t in cel.values()):
            cab = (i, cel)
            break
    if not cab:
        return {}
    i0, grupos = cab
    mapa = {}
    for linha in linhas[i0 + 1:]:
        if any(_txt(c).lower().startswith('documento') for c in linha):
            break
        for j, g in grupos.items():
            if j < len(linha) and _txt(linha[j]):
                mapa[_chave(linha[j])] = g
    return mapa


def _grupo_da_aba(titulo, mapa_indice):
    k = _chave(titulo)
    if k in mapa_indice:
        return mapa_indice[k]
    for ki, g in mapa_indice.items():  # nomes parecidos (ex.: "TESS" x "Traçadora Tess 1")
        if len(ki) >= 4 and (ki in k or k in ki):
            return g
    return ''


# --------------------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------------------

def importar_planilha(arquivo, nome_arquivo=None, setor=None, codigo=None, revisao=None, titulo=None,
                      usuario=None, publicar=False, substituir=False):
    """Importa a estrutura de um formulário. Retorna um dict com o resumo.

    `arquivo`: caminho ou objeto de arquivo. Código/revisão/título vêm do nome do arquivo quando não informados.
    Se a revisão já existir: com `substituir` e ainda em rascunho, ela é recriada; senão levanta ValueError.
    """
    nome_arquivo = nome_arquivo or (arquivo if isinstance(arquivo, str) else getattr(arquivo, 'name', ''))
    detectado = parse_arquivo(nome_arquivo)
    codigo = (codigo or (detectado[0] if detectado else '')).strip().upper()
    revisao = (revisao or (detectado[1] if detectado else '')).strip().upper()
    titulo = (titulo or (detectado[2] if detectado else codigo)).strip()
    if not codigo or not revisao:
        raise ValueError('Não foi possível identificar o código e a revisão pelo nome do arquivo; informe-os manualmente.')

    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        mapa_indice = _grupos_do_indice(wb)
        abas, ignoradas, ocultas = [], [], []
        for ordem, aba in enumerate(wb.sheetnames, 1):
            if _chave(aba) in ABAS_IGNORADAS:
                continue
            if wb[aba].sheet_state != 'visible':  # aba oculta (hidden/veryHidden) não é importada
                ocultas.append(aba)
                continue
            if 'porhora' in _chave(aba):
                ignoradas.append((aba, 'registro horário (fora do escopo desta importação)'))
                continue
            linhas = list(wb[aba].iter_rows(min_row=1, max_row=MAX_LINHAS, max_col=MAX_COLUNAS, values_only=True))
            cab = _achar_cabecalho(linhas)
            if not cab:
                ignoradas.append((aba, 'cabeçalho não reconhecido'))
                continue
            i_cab, layout, m = cab
            campos = _linhas_campos(linhas, i_cab, layout, m)
            if not campos:
                ignoradas.append((aba, 'sem itens'))
                continue
            abas.append((ordem, aba, layout, campos))
    finally:
        wb.close()

    with transaction.atomic():
        if setor is None:
            raise ValueError('Informe o setor do formulário.')
        formulario, criado = Formulario.objects.get_or_create(codigo=codigo, defaults={'setor': setor, 'titulo': titulo})
        existente = formulario.versoes.filter(revisao=revisao).first()
        if existente:
            if not substituir:
                raise ValueError(f'A revisão {codigo}.{revisao} já existe.')
            if not existente.editavel:
                raise ValueError(f'A revisão {codigo}.{revisao} já está {existente.get_status_display().lower()} e não pode ser substituída.')
            existente.delete()
        versao = FormularioVersao.objects.create(
            formulario=formulario, revisao=revisao, criado_por=usuario,
            motivo=f'Importada da planilha {nome_arquivo}'[:500])

        cad = _Cadastros()
        # grupos ajustados à mão nas revisões anteriores valem para a mesma página na nova revisão
        herdados = {}
        for nome_p, grupo_p in (Pagina.objects.filter(versao__formulario=formulario).exclude(grupo='')
                                .order_by('versao__criado_em').values_list('nome', 'grupo')):
            herdados[nome_p] = grupo_p
        total_campos = 0
        paginas_resumo = []
        for ordem, aba, layout, campos in abas:
            titulo_aba = ' '.join(aba.split())
            m_prefixo = re.match(r'^(Par[aâ]metro de \w+)\s*-\s*(.+)$', titulo_aba, re.I)
            if m_prefixo:
                grupo = m_prefixo.group(1)
                nome_pagina = f"{m_prefixo.group(2)} - {grupo.split(' de ', 1)[-1]}"
            else:
                grupo, nome_pagina = _grupo_da_aba(titulo_aba, mapa_indice), titulo_aba
                if not grupo and layout == 'etapa':
                    grupo = 'Controle de processo'
            grupo = herdados.get(nome_pagina[:120]) or grupo
            if grupo:
                formulario.garantir_grupo(grupo[:80])
            pagina = Pagina.objects.create(versao=versao, nome=nome_pagina[:120], codigo_sistema=aba.strip()[:60],
                                           grupo=grupo[:80], ordem=ordem)
            objetos = []
            for n, c in enumerate(campos, 1):
                tipo, lista = _tipo_e_limites(c, cad)
                freq, freq_detalhe = cad.frequencia(c['frequencia'])
                objetos.append(Campo(
                    pagina=pagina, item=c['item'][:20], etapa=c['etapa'][:250], descricao=c['descricao'][:300],
                    especificacao=c['especificacao'], metodo=c['metodo'][:300], plano_reacao=c['plano_reacao'],
                    responsavel=c['responsavel'][:60], frequencia=freq, frequencia_detalhe=freq_detalhe, tipo=tipo, lista=lista,
                    minimo=c['minimo'], maximo=c['maximo'], unidade=c['unidade'][:20], ordem=n))
            Campo.objects.bulk_create(objetos)
            total_campos += len(objetos)
            paginas_resumo.append((nome_pagina, grupo, layout, len(objetos)))
        if publicar:
            versao.publicar()

    return {
        'formulario': formulario, 'versao': versao, 'formulario_criado': criado, 'paginas': paginas_resumo,
        'total_campos': total_campos, 'ignoradas': ignoradas, 'ocultas': ocultas, 'frequencias_criadas': cad.criadas_freq,
    }


def setor_por_nome(nome):
    """Setor existente (sem diferenciar maiúsculas) ou novo."""
    return Setor.objects.filter(nome__iexact=nome).first() or Setor.objects.create(nome=nome)

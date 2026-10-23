"""Exportação dos registros em planilha (.xlsx), no formato das planilhas de PCP: itens nas linhas, dia x turno nas colunas.

Turnos em que a máquina foi marcada como "Em Manutenção" ou "Parada" saem com TODOS os itens preenchidos com essa situação,
em vez de ficarem em branco. Uma aba de ocorrências e uma aba com as situações da máquina acompanham as páginas.
"""
import calendar
import datetime
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pcp.models import Ocorrencia, Resposta, Turno, Verificacao
from pcp.pendencias import chave, slots_parados

ROTULOS = {v: n for v, n in Verificacao.STATUS_MAQUINA if v}
COR_STATUS = {Verificacao.MANUTENCAO: 'FFE0B2', Verificacao.PARADA: 'D9D9D9'}
COLUNAS_FIXAS = ['Nº', 'Item', 'O que verificar', 'Especificação', 'Método', 'Frequência', 'Plano de reação']
LARGURAS_FIXAS = [6, 9, 42, 36, 24, 18, 30]
NEGRITO = Font(bold=True)
CABECALHO = PatternFill('solid', fgColor='14284B')
BRANCO = Font(bold=True, color='FFFFFF')


def _nome_aba(nome, usados):
    base = re.sub(r'[\[\]:*?/\\]', '-', nome).strip()[:28] or 'Pagina'
    candidato, n = base, 2
    while candidato.lower() in usados:
        candidato = f'{base[:26]}_{n}'
        n += 1
    usados.add(candidato.lower())
    return candidato


def _valor_celula(campo, valor):
    if campo.tipo == campo.NUMERO:
        try:
            return float(Decimal(valor.replace(',', '.')))
        except InvalidOperation:
            return valor
    return valor


def periodo_do_mes(ano, mes):
    ultimo = calendar.monthrange(ano, mes)[1]
    return datetime.date(ano, mes, 1), datetime.date(ano, mes, ultimo)


def _aba_pagina(wb, formulario, versao, pagina, ini, fim, turnos, usados):
    ws = wb.create_sheet(_nome_aba(pagina.nome, usados))
    dias = [ini + datetime.timedelta(days=n) for n in range((fim - ini).days + 1)]
    ws['A1'] = f'{formulario.codigo} · {versao.revisao} — {pagina.nome} — {ini:%m/%Y}'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = ('Turnos em "Em Manutenção" ou "Parada": todos os itens saem preenchidos com a situação da máquina '
                '(laranja = manutenção, cinza = parada).')
    ws['A2'].font = Font(italic=True, color='666666')
    for j, titulo in enumerate(COLUNAS_FIXAS, 1):
        for linha in (4, 5):
            c = ws.cell(row=linha, column=j, value=titulo if linha == 4 else None)
            c.font, c.fill = BRANCO, CABECALHO
        ws.column_dimensions[get_column_letter(j)].width = LARGURAS_FIXAS[j - 1]
    col0 = len(COLUNAS_FIXAS) + 1
    colunas = {}  # (dia, turno_id) -> coluna
    for i, dia in enumerate(dias):
        primeira = col0 + i * len(turnos)
        ws.merge_cells(start_row=4, start_column=primeira, end_row=4, end_column=primeira + len(turnos) - 1)
        cab = ws.cell(row=4, column=primeira, value=dia.strftime('%d/%m'))
        cab.font, cab.fill, cab.alignment = BRANCO, CABECALHO, Alignment(horizontal='center')
        for k, t in enumerate(turnos):
            c = ws.cell(row=5, column=primeira + k, value=t.sigla)
            c.font, c.fill, c.alignment = BRANCO, CABECALHO, Alignment(horizontal='center')
            ws.column_dimensions[get_column_letter(primeira + k)].width = 9
            colunas[(dia, t.pk)] = primeira + k

    # lançamentos do mês (das páginas de mesmo nome em qualquer revisão) e situação da máquina
    respostas = defaultdict(dict)
    qs = (Resposta.objects.filter(verificacao__pagina__versao__formulario=formulario, verificacao__pagina__nome=pagina.nome,
                                  verificacao__data__range=(ini, fim)).exclude(valor='')
          .values_list('campo__item', 'campo__etapa', 'campo__descricao', 'verificacao__data', 'verificacao__turno_id', 'valor', 'conforme'))
    for item, etapa, desc, dia, turno_id, valor, conforme in qs:
        respostas[chave(item, etapa, desc)][(dia, turno_id)] = (valor, conforme)
    parados = slots_parados(pagina, ini, fim)

    for linha, campo in enumerate(pagina.campos.filter(ativo=True).select_related('frequencia'), start=6):
        fixos = [campo.ordem, campo.item, campo.descricao, campo.especificacao, campo.metodo,
                 str(campo.frequencia or ''), campo.plano_reacao]
        for j, valor in enumerate(fixos, 1):
            ws.cell(row=linha, column=j, value=valor).alignment = Alignment(wrap_text=True, vertical='top')
        if campo.critico:
            ws.cell(row=linha, column=2).font = Font(bold=True, color='C00000')
        lancados = respostas.get(chave(campo.item, campo.etapa, campo.descricao), {})
        for (dia, turno_id), coluna in colunas.items():
            status = parados.get((dia, turno_id))
            celula = ws.cell(row=linha, column=coluna)
            if status:
                celula.value = ROTULOS[status]
                celula.fill = PatternFill('solid', fgColor=COR_STATUS[status])
                celula.alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
            elif (dia, turno_id) in lancados:
                valor, conforme = lancados[(dia, turno_id)]
                celula.value = _valor_celula(campo, valor)
                celula.alignment = Alignment(horizontal='center')
                if conforme is False:
                    celula.font = Font(color='C00000', bold=True)
    ws.freeze_panes = ws.cell(row=6, column=col0)
    return ws


def _aba_situacoes(wb, formulario, paginas, ini, fim, turnos, usados):
    ws = wb.create_sheet(_nome_aba('Situação da máquina', usados))
    ws.append(['Página', 'Data', 'Turno', 'Situação', 'Motivo'])
    for c in ws[1]:
        c.font, c.fill = BRANCO, CABECALHO
    nomes = {p.nome for p in paginas}
    qs = (Verificacao.objects.filter(pagina__versao__formulario=formulario, pagina__nome__in=nomes, data__range=(ini, fim))
          .exclude(status='').select_related('turno').order_by('pagina__nome', 'data', 'turno__ordem').distinct())
    for v in qs:
        ws.append([v.pagina.nome, v.data, v.turno.sigla, ROTULOS[v.status], v.motivo])
    for coluna, largura in zip('ABCDE', (34, 12, 8, 16, 40)):
        ws.column_dimensions[coluna].width = largura
    for linha in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        linha[0].number_format = 'dd/mm/yyyy'
    return ws


def _aba_ocorrencias(wb, formulario, paginas, ini, fim, usados):
    ws = wb.create_sheet(_nome_aba('Ocorrências', usados))
    ws.append(['Data', 'Turno', 'Página', 'Item', 'Tipo', 'Descrição', 'Situação', 'Tratativa'])
    for c in ws[1]:
        c.font, c.fill = BRANCO, CABECALHO
    nomes = {p.nome for p in paginas}
    qs = (Ocorrencia.objects.filter(verificacao__pagina__versao__formulario=formulario, verificacao__pagina__nome__in=nomes,
                                    verificacao__data__range=(ini, fim))
          .select_related('verificacao__pagina', 'verificacao__turno', 'campo').order_by('verificacao__data', 'verificacao__turno__ordem'))
    for o in qs:
        situacao = ('Em aberto' if o.status == Ocorrencia.ABERTA else 'Tratada') if o.tipo == Ocorrencia.TRATATIVA else 'Observação'
        ws.append([o.verificacao.data, o.verificacao.turno.sigla, o.verificacao.pagina.nome, o.campo.item if o.campo else '',
                   o.get_tipo_display(), o.descricao, situacao, o.tratativa])
    for coluna, largura in zip('ABCDEFGH', (12, 8, 30, 10, 26, 50, 12, 40)):
        ws.column_dimensions[coluna].width = largura
    for linha in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        linha[0].number_format = 'dd/mm/yyyy'
    return ws


def montar_planilha(formulario, versao, paginas, ano, mes):
    """Workbook com uma aba por página (itens x dia/turno do mês), mais as abas de situação da máquina e de ocorrências."""
    turnos = list(Turno.objects.filter(ativo=True))
    ini, fim = periodo_do_mes(ano, mes)
    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    paginas = list(paginas)
    for pagina in paginas:
        _aba_pagina(wb, formulario, versao, pagina, ini, fim, turnos, usados)
    _aba_situacoes(wb, formulario, paginas, ini, fim, turnos, usados)
    _aba_ocorrencias(wb, formulario, paginas, ini, fim, usados)
    return wb


# --------------------------------------------------------------------------------------
# Exportação "DADOS": tabela longa (uma linha por item x data x turno)
# --------------------------------------------------------------------------------------
COLUNAS_DADOS = ['Formulário', 'Grupo', 'Página', 'ORD', 'Item', 'Data', 'Turno', 'Valor']
LIMITE_LINHAS_DADOS = 250000


def lancamentos_da_pagina(formulario, pagina, ini, fim):
    """({chave do item: {(data, turno_id): (valor, conforme)}}, {(data, turno_id): status}) do período, juntando as revisões."""
    respostas = defaultdict(dict)
    qs = (Resposta.objects.filter(verificacao__pagina__versao__formulario=formulario, verificacao__pagina__nome=pagina.nome,
                                  verificacao__data__range=(ini, fim)).exclude(valor='')
          .values_list('campo__item', 'campo__etapa', 'campo__descricao', 'verificacao__data', 'verificacao__turno_id', 'valor', 'conforme'))
    for item, etapa, desc, dia, turno_id, valor, conforme in qs:
        respostas[chave(item, etapa, desc)][(dia, turno_id)] = (valor, conforme)
    return respostas, slots_parados(pagina, ini, fim)


def estimar_linhas_dados(paginas, ini, fim, n_turnos):
    dias = (fim - ini).days + 1
    return sum(p.campos.filter(ativo=True).count() for p in paginas) * dias * n_turnos


def montar_dados(formulario, versao, paginas, ini, fim, somente_com_valor=False):
    """Workbook com a aba EXPORTAÇÃO no formato DADOS: Formulário | Grupo | Página | ORD | Item | Data | Turno | Valor.

    ORD é o Nº do item no bloco. Uma linha por item ativo, data e turno (mesmo sem valor, a menos que `somente_com_valor`);
    turnos em manutenção/parada trazem a situação da máquina no Valor. Ordem: página (grupo e ordem), data, turno, ORD."""
    turnos = list(Turno.objects.filter(ativo=True))
    dias = [ini + datetime.timedelta(days=n) for n in range((fim - ini).days + 1)]
    wb = Workbook()
    ws = wb.active
    ws.title = 'EXPORTAÇÃO'
    ws.append(COLUNAS_DADOS)
    for c in ws[1]:
        c.font = NEGRITO
    for pagina in paginas:
        campos = list(pagina.campos.filter(ativo=True).order_by('ordem', 'pk'))
        respostas, parados = lancamentos_da_pagina(formulario, pagina, ini, fim)
        for dia in dias:
            for turno in turnos:
                status = parados.get((dia, turno.pk))
                for campo in campos:
                    if status:
                        valor = ROTULOS[status]
                    else:
                        lancado = respostas.get(chave(campo.item, campo.etapa, campo.descricao), {}).get((dia, turno.pk))
                        valor = _valor_celula(campo, lancado[0]) if lancado else None
                    if somente_com_valor and valor in (None, ''):
                        continue
                    ws.append([formulario.codigo, pagina.grupo, pagina.nome, campo.ordem, campo.item, dia, turno.sigla, valor])
    for coluna, largura in zip('ABCDEFGH', (11, 22, 24, 6, 8, 11, 7, 18)):
        ws.column_dimensions[coluna].width = largura
    for linha in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        linha[0].number_format = 'dd/mm/yyyy'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return wb

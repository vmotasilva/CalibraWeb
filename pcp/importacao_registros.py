"""Importação dos registros (lançamentos) já feitos nas planilhas, separada da importação da estrutura.

Só preenche páginas e itens que JÁ existem na revisão escolhida (importe a estrutura antes). Cada célula com valor, na
coluna de uma data e turno do ano escolhido, vira uma Resposta de uma Verificação (página + data + turno):

- abas com colunas de data (cabeçalho com a data e T1/T2/T3 embaixo): um lançamento por data e turno;
- abas com colunas de mês (Janeiro, Fevereiro...): lançamento mensal, gravado no dia 1º do mês, no primeiro turno;
- abas ocultas e abas "Por Hora" são ignoradas.

Reimportar é seguro: respostas que já existem (importadas antes ou lançadas no sistema) não são sobrescritas.
"""
import datetime
from collections import Counter, defaultdict

import openpyxl
from django.db import transaction

from pcp.importacao import (ABAS_IGNORADAS, MAX_LINHAS, RE_NUM, _achar_cabecalho, _chave, _linhas_campos, _txt,
                            nome_da_pagina, parse_arquivo)
from pcp.models import Formulario, FormularioVersao, ListaOpcoes, Pagina, Resposta, Turno, Verificacao

MESES = {'janeiro': 1, 'fevereiro': 2, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7, 'agosto': 8,
         'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12}
PRIMEIRA_COLUNA_DE_DADOS = 5  # as colunas de estrutura (item, descrição, especificação...) ficam antes
SINONIMOS_NAO_REALIZADO = {'nao', 'falta', 'falt'}  # 'não', 'falta' = item não realizado
IGNORAR_VALOR = {'', '-', '--', '—', 'n/a', 'na'}  # sem lançamento (N/A é "não se aplica", não um resultado)


def _como_data(valor):
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    return None


def colunas_de_dados(linhas, i_cab, ano, turnos_por_chave, turno_mensal):
    """[(coluna, data, turno)] das colunas de lançamento do ano. Datas podem estar na linha do cabeçalho ou na anterior."""
    janela = range(max(i_cab - 1, 0), min(i_cab + 3, len(linhas)))
    i_turnos = next((i for i in janela if any(_chave(c) in ('t1', 't2', 't3') for c in linhas[i])), None)
    colunas = []
    for i in janela:
        for j, c in enumerate(linhas[i]):
            if j < PRIMEIRA_COLUNA_DE_DADOS:
                continue
            data = _como_data(c)
            if data and data.year == ano and i_turnos is not None:
                for k in range(3):  # a data é uma célula mesclada sobre T1/T2/T3
                    if j + k < len(linhas[i_turnos]):
                        turno = turnos_por_chave.get(_chave(linhas[i_turnos][j + k]))
                        if turno:
                            colunas.append((j + k, data, turno))
            elif isinstance(c, str) and _chave(c) in MESES and i == i_cab and turno_mensal:
                colunas.append((j, datetime.date(ano, MESES[_chave(c)], 1), turno_mensal))
    return colunas


def _normalizar(campo, valor, opcoes):
    """Ajusta o valor lido ao vocabulário da lista do campo (ok -> Ok, 'Em Manut.' -> Manut.)."""
    if campo.tipo != campo.LISTA or not opcoes:
        return valor
    k = _chave(valor)
    por_chave = {_chave(o): o for o in opcoes}
    if k in por_chave:
        return por_chave[k]
    for trecho, destino in (('manut', 'Manut.'), ('maqparada', 'Máq. parada'), ('maquinaparada', 'Máq. parada')):
        if trecho in k and destino in opcoes:
            return destino
    return valor


def importar_registros(arquivo, nome_arquivo=None, ano=2026, formulario=None, revisao=None, simular=False,
                       rastreio_desde_primeiro=False, hoje=None, ajustar_tipos=False):
    """Importa os lançamentos do ano. Retorna um resumo (dict). Com `simular`, nada é gravado.

    Lançamentos em datas posteriores a `hoje` (padrão: o dia de hoje) são ignorados e contados em `futuros`."""
    hoje = hoje or datetime.date.today()
    nome_arquivo = nome_arquivo or (arquivo if isinstance(arquivo, str) else getattr(arquivo, 'name', ''))
    detectado = parse_arquivo(nome_arquivo)
    if formulario is None:
        if not detectado:
            raise ValueError('Não foi possível identificar o formulário pelo nome do arquivo; informe-o.')
        formulario = Formulario.objects.filter(codigo=detectado[0]).first()
        if formulario is None:
            raise ValueError(f'O formulário {detectado[0]} não existe. Importe a estrutura antes de importar os registros.')
    rev = (revisao or (detectado[1] if detectado else '')).strip().upper()
    versao = (formulario.versoes.filter(revisao=rev).first() if rev else None) or formulario.versao_vigente
    if versao is None:
        raise ValueError(f'O formulário {formulario.codigo} não tem a revisão {rev or "vigente"}. Importe a estrutura antes.')

    turnos = list(Turno.objects.filter(ativo=True))
    if not turnos:
        raise ValueError('Cadastre ao menos um turno ativo (T1, T2, T3) antes de importar registros.')
    turnos_por_chave = {_chave(t.sigla): t for t in turnos}
    turno_mensal = turnos[0]
    listas = list(ListaOpcoes.objects.prefetch_related('opcoes'))
    opcoes_por_lista = {l.pk: [o.valor for o in l.opcoes.all()] for l in listas}
    conforme_por_lista = {l.pk: {o.valor: o.conforme for o in l.opcoes.all()} for l in listas}
    conforme_global = {}  # vocabulário de todas as listas: 'Realizado' num item Ok/Nok continua valendo como conforme
    for l in listas:
        for o in l.opcoes.all():
            conforme_global.setdefault(_chave(o.valor), (o.valor, o.conforme))

    resumo = {'versao': versao, 'ano': ano, 'abas': [], 'sem_pagina': [], 'ocultas': [], 'ignoradas': [], 'linhas_sem_item': 0,
              'verificacoes_novas': 0, 'respostas_novas': 0, 'respostas_existentes': 0, 'fora_do_limite': 0, 'primeira_data': None,
              'vigencia_ajustada': None, 'futuros': 0, 'nao_reconhecidos': [], 'tipos_ajustados': []}
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    lancamentos = []  # (pagina, data, turno, campo, valor)
    ignorar = {_chave(x) for x in IGNORAR_VALOR}
    try:
        for aba in wb.sheetnames:
            if _chave(aba) in ABAS_IGNORADAS or 'porhora' in _chave(aba):
                continue
            ws = wb[aba]
            if ws.sheet_state != 'visible':
                resumo['ocultas'].append(aba)
                continue
            linhas = list(ws.iter_rows(min_row=1, max_row=MAX_LINHAS, values_only=True))
            cab = _achar_cabecalho([l[:14] for l in linhas])
            if not cab:
                resumo['ignoradas'].append((aba, 'cabeçalho não reconhecido'))
                continue
            i_cab, layout, mapa = cab
            colunas = colunas_de_dados(linhas, i_cab, ano, turnos_por_chave, turno_mensal)
            if not colunas:
                resumo['ignoradas'].append((aba, f'sem colunas de {ano}'))
                continue
            pagina = versao.paginas.filter(nome=nome_da_pagina(aba)).first()
            if pagina is None:
                resumo['sem_pagina'].append(aba)
                continue
            # casa cada linha da aba com o item cadastrado (mesmo item/etapa/descrição, na ordem em caso de repetição)
            disponiveis = defaultdict(list)
            for campo in pagina.campos.select_related('lista'):
                disponiveis[(campo.item, campo.etapa, campo.descricao)].append(campo)
            n_aba = 0
            for c in _linhas_campos([l[:14] for l in linhas], i_cab, layout, mapa):
                chave = (c['item'][:20], c['etapa'][:250], c['descricao'][:300])
                if not disponiveis[chave]:
                    resumo['linhas_sem_item'] += 1
                    continue
                campo = disponiveis[chave].pop(0)
                completa = linhas[c['_indice']]  # linha inteira, com as colunas de data x turno
                opcoes = opcoes_por_lista.get(campo.lista_id, [])
                for coluna, data, turno in colunas:
                    valor = _txt(completa[coluna]) if coluna < len(completa) else ''
                    if _chave(valor) in ignorar or valor.startswith('#'):
                        continue
                    if data > hoje:
                        resumo['futuros'] += 1
                        continue
                    lancamentos.append((pagina, data, turno, campo, _normalizar(campo, valor, opcoes)[:255]))
                    n_aba += 1
            resumo['abas'].append((aba, pagina.nome, n_aba, len(colunas)))
    finally:
        wb.close()

    resumo['lancamentos'] = len(lancamentos)
    resumo['primeira_data'] = min((l[1] for l in lancamentos), default=None)
    resumo['ultima_data'] = max((l[1] for l in lancamentos), default=None)
    with transaction.atomic():
        if ajustar_tipos:
            # item de lista cujos lançamentos são quase todos numéricos passa a ser numérico (o tipo da estrutura era só uma dedução)
            por_campo, instancias = defaultdict(list), {}
            for pagina, data, turno, campo, valor in lancamentos:
                por_campo[campo.pk].append(valor)
                instancias[campo.pk] = (pagina, campo)
            for pk, valores in por_campo.items():
                pagina, campo = instancias[pk]
                numericos = sum(1 for v in valores if RE_NUM.match(v))
                if campo.tipo == campo.LISTA and len(valores) >= 5 and numericos / len(valores) >= 0.8:
                    campo.tipo, campo.lista = campo.NUMERO, None
                    campo.save(update_fields=['tipo', 'lista'])
                    resumo['tipos_ajustados'].append((pagina.nome, campo.item or campo.descricao[:40], len(valores)))
        paginas_ids = {l[0].pk for l in lancamentos}
        verificacoes = {(v.pagina_id, v.data, v.turno_id): v
                        for v in Verificacao.objects.filter(pagina_id__in=paginas_ids, data__year=ano)}
        a_criar = {}
        for pagina, data, turno, campo, valor in lancamentos:
            chave = (pagina.pk, data, turno.pk)
            if chave not in verificacoes and chave not in a_criar:
                a_criar[chave] = Verificacao(pagina=pagina, data=data, turno=turno)
        Verificacao.objects.bulk_create(a_criar.values(), batch_size=2000)
        if a_criar:  # recarrega para obter as chaves primárias (nem todo banco as devolve no bulk_create)
            verificacoes = {(v.pagina_id, v.data, v.turno_id): v
                            for v in Verificacao.objects.filter(pagina_id__in=paginas_ids, data__year=ano)}
        resumo['verificacoes_novas'] = len(a_criar)
        existentes = set(Resposta.objects.filter(verificacao__pagina_id__in=paginas_ids, verificacao__data__year=ano)
                         .values_list('verificacao_id', 'campo_id'))
        nao_reconhecidos = Counter()
        novas = []
        for pagina, data, turno, campo, valor in lancamentos:
            v = verificacoes[(pagina.pk, data, turno.pk)]
            if (v.pk, campo.pk) in existentes:
                resumo['respostas_existentes'] += 1
                continue
            existentes.add((v.pk, campo.pk))
            if campo.tipo == campo.LISTA:
                conforme = conforme_por_lista.get(campo.lista_id, {}).get(valor)
                k = _chave(valor)
                if conforme is None and k in conforme_global:
                    valor, conforme = conforme_global[k]
                elif conforme is None and (k in SINONIMOS_NAO_REALIZADO or k.startswith('naor')):
                    valor, conforme = 'Não realizado', False
                if conforme is None:
                    nao_reconhecidos[(campo.lista.nome if campo.lista_id else '-', valor)] += 1
            else:
                conforme = campo.avaliar(valor)
            resumo['fora_do_limite'] += int(conforme is False)
            novas.append(Resposta(verificacao=v, campo=campo, valor=valor, conforme=conforme))
        resumo['nao_reconhecidos'] = [(lista, valor, n) for (lista, valor), n in nao_reconhecidos.most_common(15)]
        Resposta.objects.bulk_create(novas, batch_size=2000)
        resumo['respostas_novas'] = len(novas)
        if rastreio_desde_primeiro and resumo['primeira_data']:
            # o rastreio de atrasos começa na primeira vigência do formulário: recua-a até o primeiro registro importado
            primeira = min((v.vigencia for v in formulario.versoes.all() if v.vigencia), default=None)
            if primeira is None or resumo['primeira_data'] < primeira:
                alvo = formulario.versao_vigente or versao
                alvo.vigencia = resumo['primeira_data']
                alvo.save(update_fields=['vigencia'])
                resumo['vigencia_ajustada'] = resumo['primeira_data']
        if simular:
            transaction.set_rollback(True)
    resumo['simulado'] = simular
    return resumo

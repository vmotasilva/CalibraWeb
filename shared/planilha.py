"""Subconjunto do pandas usado pelo sistema, implementado só com openpyxl/csv.

O pandas (e o numpy que ele arrasta) estourava o limite de tamanho da função na Vercel. Este módulo oferece a
parte da API que o sistema realmente usa (leitura de xlsx/csv, DataFrame simples, escrita de xlsx, notna/isna,
to_datetime), com as mesmas convenções: célula vazia = NaN, ``Timestamp`` = datetime, etc.
Use ``from shared import planilha as pd`` no lugar de ``import pandas as pd``.
"""
import csv
import datetime
import io
import math
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

nan = float('nan')
Timestamp = datetime.datetime


class _NaTType:
    """Data ausente (NaT). Falsa; date()/time() devolvem ela mesma, como o pandas."""
    def __bool__(self):
        return False

    def date(self):
        return self

    def time(self):
        return self

    def __repr__(self):
        return 'NaT'

    __str__ = __repr__


NaT = _NaTType()


class _Errors:
    class EmptyDataError(ValueError):
        pass


errors = _Errors


def isna(v):
    if v is None or v is NaT:
        return True
    if isinstance(v, float):
        return math.isnan(v)
    return False


isnull = isna


def notna(v):
    return not isna(v)


notnull = notna


# ----------------------------------------------------------------------------- datas
_FORMATOS_ISO = ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d', '%Y/%m/%d', '%Y-%m-%dT%H:%M:%S')
_FORMATOS_DMY = ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d/%m/%y', '%d-%m-%y')
_FORMATOS_MDY = ('%m/%d/%Y %H:%M:%S', '%m/%d/%Y %H:%M', '%m/%d/%Y', '%m-%d-%Y', '%m/%d/%y')


def to_datetime(valor, format=None, dayfirst=False, errors='raise'):
    """Converte texto/data em datetime. Textos vazios viram NaT. Com errors='coerce' falhas viram NaT."""
    try:
        if isinstance(valor, datetime.datetime):
            return valor
        if isinstance(valor, datetime.date):
            return datetime.datetime.combine(valor, datetime.time())
        if isna(valor):
            return NaT
        texto = str(valor).strip()
        if not texto:
            return NaT
        if format:
            return datetime.datetime.strptime(texto, format)
        ordem = (_FORMATOS_ISO + _FORMATOS_DMY + _FORMATOS_MDY) if dayfirst else (_FORMATOS_ISO + _FORMATOS_MDY + _FORMATOS_DMY)
        for fmt in ordem:
            try:
                return datetime.datetime.strptime(texto, fmt)
            except ValueError:
                continue
        raise ValueError(f'Data inválida: {texto}')
    except (ValueError, TypeError):
        if errors == 'coerce':
            return NaT
        raise


# ----------------------------------------------------------------------------- acessador .str
class _Str:
    """Operações de texto em Series/Index, elemento a elemento (valores ausentes são preservados)."""

    def __init__(self, seq, fabrica):
        self._seq, self._fab = list(seq), fabrica

    def _aplica(self, fn):
        return self._fab([v if isna(v) else fn(v) for v in self._seq])

    def strip(self):
        return self._aplica(lambda v: str(v).strip())

    def upper(self):
        return self._aplica(lambda v: str(v).upper())

    def lower(self):
        return self._aplica(lambda v: str(v).lower())

    def replace(self, pat, repl, regex=False, **_):
        if regex:
            rx = re.compile(pat)
            return self._aplica(lambda v: rx.sub(repl, str(v)))
        return self._aplica(lambda v: str(v).replace(pat, repl))

    def normalize(self, forma):
        return self._aplica(lambda v: unicodedata.normalize(forma, str(v)))

    def encode(self, codificacao, errors='strict'):
        return self._aplica(lambda v: str(v).encode(codificacao, errors=errors))

    def decode(self, codificacao, errors='strict'):
        return self._aplica(lambda v: v.decode(codificacao, errors=errors) if isinstance(v, bytes) else v)

    def len(self):
        return self._aplica(lambda v: len(str(v)))


class Index(list):
    @property
    def str(self):
        return _Str(self, Index)

    def tolist(self):
        return list(self)


class Series:
    def __init__(self, valores=(), nome=None):
        self._v, self.name = list(valores), nome

    def __iter__(self):
        return iter(self._v)

    def __len__(self):
        return len(self._v)

    def __getitem__(self, i):
        return self._v[i]

    @property
    def empty(self):
        return not self._v

    @property
    def iloc(self):
        return _Iloc(self._v)

    @property
    def str(self):
        return _Str(self._v, Series)

    def tolist(self):
        return list(self._v)

    def dropna(self):
        return Series([v for v in self._v if not isna(v)], self.name)

    def astype(self, tipo):
        if tipo is str:
            return Series([('nan' if isna(v) else str(v)) for v in self._v], self.name)
        return Series([(tipo(v) if not isna(v) else v) for v in self._v], self.name)

    def map(self, fn):
        return Series([fn(v) for v in self._v], self.name)

    def max(self):
        validos = [v for v in self._v if not isna(v)]
        return max(validos) if validos else nan

    def fillna(self, valor):
        return Series([valor if isna(v) else v for v in self._v], self.name)


class _Iloc:
    def __init__(self, seq):
        self._s = seq

    def __getitem__(self, i):
        return self._s[i]


class Row(dict):
    """Linha de DataFrame (como pandas.Series): get() devolve o padrão só quando a coluna não existe."""

    def __getitem__(self, chave):
        return dict.__getitem__(self, chave)


# ----------------------------------------------------------------------------- DataFrame
class DataFrame:
    def __init__(self, data=None, columns=None):
        if isinstance(data, DataFrame):
            cols, linhas = list(data.columns), [list(r) for r in data._linhas]
            if columns is not None:
                cols, linhas = list(columns), [[r[data.columns.index(c)] if c in data.columns else nan for c in columns] for r in linhas]
        elif isinstance(data, dict):
            cols = list(columns) if columns is not None else list(data.keys())
            tamanho = max((len(v) if isinstance(v, (list, tuple, Series)) else 1 for v in data.values()), default=0)
            linhas = []
            for i in range(tamanho):
                linha = []
                for c in cols:
                    v = data.get(c, [])
                    if isinstance(v, (list, tuple, Series)):
                        linha.append(list(v)[i] if i < len(v) else nan)
                    else:
                        linha.append(v)
                linhas.append(linha)
        elif data is None:
            cols, linhas = list(columns or []), []
        else:
            data = list(data)
            if data and isinstance(data[0], dict):
                cols = list(columns) if columns is not None else []
                if columns is None:
                    for d in data:
                        for k in d:
                            if k not in cols:
                                cols.append(k)
                linhas = [[d.get(c, nan) for c in cols] for d in data]
            else:
                cols = list(columns) if columns is not None else list(range(len(data[0]) if data else 0))
                linhas = [list(r) for r in data]
        self._cols = Index(cols)
        self._linhas = linhas

    # -- estrutura
    @property
    def columns(self):
        return self._cols

    @columns.setter
    def columns(self, novas):
        novas = list(novas)
        if len(novas) != len(self._cols):
            raise ValueError('Quantidade de colunas diferente da existente')
        self._cols = Index(novas)

    @property
    def empty(self):
        return not self._linhas or not self._cols

    @property
    def shape(self):
        return (len(self._linhas), len(self._cols))

    def __len__(self):
        return len(self._linhas)

    def __contains__(self, coluna):
        return coluna in self._cols

    def __iter__(self):
        return iter(self._cols)

    def _idx(self, coluna):
        try:
            return self._cols.index(coluna)
        except ValueError:
            raise KeyError(coluna)

    def __getitem__(self, chave):
        if isinstance(chave, (list, tuple)):
            idx = [self._idx(c) for c in chave]
            return DataFrame([[r[i] for i in idx] for r in self._linhas], columns=list(chave))
        i = self._idx(chave)
        return Series([r[i] for r in self._linhas], chave)

    def __setitem__(self, coluna, valor):
        n = len(self._linhas)
        if isinstance(valor, (list, tuple, Series)):
            valores = list(valor)
        else:
            valores = [valor] * n
        if coluna in self._cols:
            i = self._idx(coluna)
            for r, v in zip(self._linhas, valores):
                r[i] = v
        else:
            self._cols.append(coluna)
            for r, v in zip(self._linhas, valores):
                r.append(v)

    @property
    def iloc(self):
        return _IlocDF(self)

    def copy(self):
        return DataFrame(self)

    # -- conteúdo
    def iterrows(self):
        for n, r in enumerate(self._linhas):
            yield n, Row(zip(self._cols, r))

    def rename(self, columns=None, inplace=False):
        mapa = columns or {}
        novas = [mapa.get(c, c) for c in self._cols]
        if inplace:
            self._cols = Index(novas)
            return None
        out = DataFrame(self)
        out._cols = Index(novas)
        return out

    def fillna(self, valor):
        out = DataFrame(self)
        out._linhas = [[valor if isna(v) else v for v in r] for r in out._linhas]
        return out

    def drop_duplicates(self, subset=None):
        idx = [self._idx(c) for c in (subset or self._cols)]
        vistos, linhas = set(), []
        for r in self._linhas:
            chave = tuple(r[i] for i in idx)
            if chave not in vistos:
                vistos.add(chave)
                linhas.append(r)
        out = DataFrame(self)
        out._linhas = linhas
        return out

    def groupby(self, chaves):
        chaves = [chaves] if isinstance(chaves, str) else list(chaves)
        idx = [self._idx(c) for c in chaves]
        grupos = {}
        for r in self._linhas:
            grupos.setdefault(tuple(r[i] for i in idx), []).append(r)
        try:
            ordem = sorted(grupos)
        except TypeError:
            ordem = list(grupos)
        out = []
        for k in ordem:
            sub = DataFrame(self)
            sub._linhas = grupos[k]
            out.append((k if len(k) > 1 else k[0], sub))
        return out

    def to_excel(self, destino, sheet_name='Sheet1', index=False, engine=None):
        if isinstance(destino, ExcelWriter):
            destino._adicionar(self, sheet_name)
            return
        with ExcelWriter(destino) as w:
            w._adicionar(self, sheet_name)


class _IlocDF:
    def __init__(self, df):
        self._df = df

    def __getitem__(self, i):
        return Row(zip(self._df._cols, self._df._linhas[i]))


# ----------------------------------------------------------------------------- escrita
def _celula(v):
    if isna(v):
        return None
    return v


class ExcelWriter:
    """with ExcelWriter(destino) as w: df.to_excel(w, sheet_name=...). Destino: caminho, BytesIO ou resposta HTTP."""

    def __init__(self, destino, engine=None):
        self._destino = destino
        self.book = Workbook()
        self.book.remove(self.book.active)
        self.sheets = {}

    def _adicionar(self, df, nome):
        ws = self.book.create_sheet(str(nome)[:31])
        ws.append([c for c in df.columns])
        for c in ws[1]:
            c.font = Font(bold=True)
        for linha in df._linhas:
            ws.append([_celula(v) for v in linha])
        self.sheets[str(nome)[:31]] = ws

    def close(self):
        if not self.book.worksheets:
            self.book.create_sheet('Sheet1')
        self.book.save(self._destino)

    def __enter__(self):
        return self

    def __exit__(self, tipo, valor, tb):
        if tipo is None:
            self.close()


# ----------------------------------------------------------------------------- leitura
def _texto_dtype(v):
    return v if isna(v) or isinstance(v, str) else str(v)


def _normaliza_celula(v):
    if v is None:
        return nan
    if isinstance(v, str):
        return v if v != '' else nan
    return v


def _cabecalho(linha):
    nomes, usados = [], {}
    for i, c in enumerate(linha):
        nome = 'Unnamed: %d' % i if c is None or str(c).strip() == '' else str(c) if not isinstance(c, str) else c
        if nome in usados:
            usados[nome] += 1
            nome = f'{nome}.{usados[nome]}'
        else:
            usados[nome] = 0
        nomes.append(nome)
    return nomes


def _df_de_linhas(linhas, dtype=None):
    linhas = [list(r) for r in linhas]
    if not linhas:
        raise errors.EmptyDataError('No columns to parse from file')
    cab = _cabecalho(linhas[0])
    corpo = [[_normaliza_celula(v) for v in r] + [nan] * (len(cab) - len(r)) for r in linhas[1:]]
    corpo = [r[:len(cab)] for r in corpo]
    while corpo and all(isna(v) for v in corpo[-1]):
        corpo.pop()  # linhas totalmente vazias no fim não contam
    df = DataFrame(corpo, columns=cab)
    if dtype is str or dtype == 'str':
        df._linhas = [[_texto_dtype(v) for v in r] for r in df._linhas]
    elif isinstance(dtype, dict):
        for c, t in dtype.items():
            if c in df.columns and (t is str or t == 'str'):
                i = df._idx(c)
                for r in df._linhas:
                    r[i] = _texto_dtype(r[i])
    return df


class ExcelFile:
    def __init__(self, fonte, engine=None):
        if engine == 'xlrd':
            raise ImportError('Leitura de .xls (xlrd) não está disponível; salve o arquivo como .xlsx.')
        if hasattr(fonte, 'seek'):
            fonte.seek(0)
        self._wb = load_workbook(fonte, read_only=True, data_only=True)
        self.sheet_names = list(self._wb.sheetnames)

    def parse(self, sheet_name=0, dtype=None, **_):
        if isinstance(sheet_name, int):
            sheet_name = self.sheet_names[sheet_name]
        elif sheet_name not in self.sheet_names:
            raise ValueError(f"Worksheet named '{sheet_name}' not found")
        return _df_de_linhas(self._wb[sheet_name].iter_rows(values_only=True), dtype)


def read_excel(fonte, sheet_name=0, dtype=None, engine=None, **_):
    xl = fonte if isinstance(fonte, ExcelFile) else ExcelFile(fonte, engine=engine)
    return xl.parse(sheet_name, dtype=dtype)


def _converte_coluna(valores):
    """Como o pandas: coluna toda numérica vira número (com vazios, float); senão fica texto."""
    pres = [v for v in valores if not isna(v)]
    if not pres:
        return valores
    try:
        if all(re.fullmatch(r'[+-]?\d+', str(v).strip()) for v in pres):
            if len(pres) == len(valores):
                return [int(str(v).strip()) for v in valores]
            return [nan if isna(v) else float(str(v).strip()) for v in valores]
        nums = [float(str(v).strip()) for v in pres]
        it = iter(nums)
        return [nan if isna(v) else next(it) for v in valores]
    except ValueError:
        return valores


def read_csv(fonte, sep=',', encoding=None, dtype=None, engine=None, **_):
    conteudo = fonte.read() if hasattr(fonte, 'read') else open(fonte, 'rb').read()
    if isinstance(conteudo, bytes):
        texto = conteudo.decode(encoding or 'utf-8')
    else:
        texto = conteudo
    texto = texto.lstrip('﻿')
    if not texto.strip():
        raise errors.EmptyDataError('No columns to parse from file')
    if sep is None:
        try:
            sep = csv.Sniffer().sniff(texto[:4096], delimiters=',;\t|').delimiter
        except csv.Error:
            sep = ','
    linhas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    df = _df_de_linhas(linhas, dtype)
    if dtype is None:
        for i in range(len(df.columns)):
            col = _converte_coluna([r[i] for r in df._linhas])
            for r, v in zip(df._linhas, col):
                r[i] = v
    return df
